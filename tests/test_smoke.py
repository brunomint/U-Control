import os
from werkzeug.security import generate_password_hash

import main as ucontrol_main
from models import DisplayState, User


def _operator_socket(app):
    """Cliente Socket.IO de um operador logado — os eventos de controle
    do telão exigem login (extensions.OperatorSocketIO)."""
    with app.app_context():
        if not User.query.filter_by(username="tester").first():
            ucontrol_main.db.session.add(User(
                username="tester",
                password=generate_password_hash("senha123", method="pbkdf2:sha256"),
            ))
            ucontrol_main.db.session.commit()
    http = app.test_client()
    http.post("/login", data={"username": "tester", "password": "senha123"})
    return ucontrol_main.socketio.test_client(app, flask_test_client=http)


def test_create_app_uses_in_memory_db_when_testing(app):
    assert app.config["TESTING"] is True
    assert app.config["SQLALCHEMY_DATABASE_URI"] == "sqlite:///:memory:"


def test_root_redirects_to_login(client):
    response = client.get("/")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_login_page_loads(client):
    response = client.get("/login")
    assert response.status_code == 200


def test_data_pt_formats_date_in_portuguese():
    assert ucontrol_main.data_pt("2026-07-08") == "Quarta-feira, 8 de julho de 2026"


def test_data_pt_returns_input_unchanged_on_invalid_date():
    assert ucontrol_main.data_pt("not-a-date") == "not-a-date"


def test_display_state_table_has_media_type_column(app):
    with app.app_context():
        columns = [c.name for c in DisplayState.__table__.columns]
    assert "media_type" in columns


def test_countdown_page_route_resolves_to_blueprint(client):
    response = client.get("/countdown")
    assert response.status_code == 200


def test_start_countdown_broadcasts_and_updates_state(app):
    socket_client = _operator_socket(app)
    socket_client.emit("start_countdown", {"duration_seconds": 60, "label": "Sermão"})
    received = socket_client.get_received()
    events = [msg["name"] for msg in received]
    assert "countdown_start" in events

    state_response = app.test_client().get("/api/countdown/state")
    assert state_response.get_json()["active"] is True
    socket_client.disconnect()


def test_set_ticker_broadcasts_and_updates_state(app):
    # o aviso vai pra sala de cada tela (public/monitor/obs), não pra todo
    # mundo — a tela se identifica pra entrar na sala
    tela = ucontrol_main.socketio.test_client(app)
    tela.emit("identify", {"role": "public"})
    socket_client = _operator_socket(app)
    socket_client.emit("set_ticker", {"text": "Bem-vindos!", "target": "all"})
    events = [msg["name"] for msg in tela.get_received()]
    assert "ticker_update" in events
    tela.disconnect()

    ticker_response = app.test_client().get("/api/ticker")
    data = ticker_response.get_json()
    assert data["text"] == "Bem-vindos!"
    assert data["active"] is True
    socket_client.disconnect()


def test_init_countdown_and_ticker_state_hydrate_from_db(app):
    # Simula o que create_app() faz no startup real: lê o rótulo/texto
    # salvos no banco pro estado em memória. Não pode rodar no import do
    # módulo (blueprints/countdown.py, ticker.py) — precisa de app_context.
    from blueprints.countdown import _countdown_state, _save_countdown_config, init_countdown_state
    from blueprints.ticker import _save_ticker_config, init_ticker_state
    import blueprints.ticker as ticker_module

    with app.app_context():
        _save_countdown_config({"label": "Rótulo salvo"})
        _save_ticker_config("Texto salvo")

        init_countdown_state()
        init_ticker_state()

    assert _countdown_state["label"] == "Rótulo salvo"
    assert ticker_module._ticker_text == "Texto salvo"


def test_login_redirects_to_dashboard_and_authenticates(app, client):
    with app.app_context():
        if not User.query.filter_by(username="tester").first():
            user = User(
                username="tester",
                password=generate_password_hash("senha123", method="pbkdf2:sha256"),
            )
            ucontrol_main.db.session.add(user)
            ucontrol_main.db.session.commit()

    response = client.post(
        "/login", data={"username": "tester", "password": "senha123"}
    )
    assert response.status_code == 302
    assert "/dashboard" in response.headers["Location"]

    # /dashboard exige sessão autenticada; client (sem login explícito)
    # segue com o cookie de sessão que o POST acima já deixou setado.
    dashboard_response = client.get("/dashboard")
    assert dashboard_response.status_code == 200


def test_login_with_wrong_password_shows_flash_message(client):
    response = client.post(
        "/login",
        data={"username": "nao-existe", "password": "errada"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "incorretos".encode() in response.data


def test_bible_verse_lookup_returns_joao_3_16(logged_in_client):
    response = logged_in_client.get(
        "/api/bible/verse",
        query_string={"book": "jo", "chapter": 3, "verse_start": 16},
    )
    data = response.get_json()
    assert "Deus amou" in data["verses"][0]["text"]
    assert data["reference"] == "João 3:16"


def test_bible_search_by_reference_resolves_to_joao_3_16(logged_in_client):
    response = logged_in_client.get("/api/bible/search", query_string={"q": "Jo 3:16"})
    data = response.get_json()
    assert data["type"] == "reference"
    assert any("Deus amou" in r["text"] for r in data["results"])


def test_update_bible_broadcasts_and_appends_to_display_log(app):
    socket_client = _operator_socket(app)
    socket_client.emit("update_bible", {"reference": "João 3:16", "text": "..."})
    received = socket_client.get_received()
    events = [msg["name"] for msg in received]
    assert "update_bible" in events
    socket_client.disconnect()


def test_bible_add_and_remove_from_order_target_todays_order(app, logged_in_client):
    # Fase 5.6: unificado na tabela Order (por data) — os botões de bíblia
    # não têm seletor de data, então miram sempre a ordem de hoje.
    import orders_store
    from datetime import date

    today = date.today().isoformat()

    with app.app_context():
        existing = orders_store.load_orders()

    add_response = logged_in_client.post(
        "/api/bible/add_to_order", json={"reference": "João 3:16"}
    )
    add_json = add_response.get_json()
    assert add_json["success"] is True
    assert add_json["date"] == today
    assert "hoje" in add_json["message"]

    with app.app_context():
        order = orders_store.load_orders().get(today, {})
        itens = order.get("itens") or []
        assert any(i.get("isBible") and i.get("reference") == "João 3:16" for i in itens)

    remove_response = logged_in_client.post(
        "/api/bible/remove_from_order", json={"reference": "João 3:16"}
    )
    assert remove_response.get_json()["success"] is True

    with app.app_context():
        order = orders_store.load_orders().get(today, {})
        itens = order.get("itens") or []
        assert not any(i.get("isBible") and i.get("reference") == "João 3:16" for i in itens)
        # limpeza — banco de teste é compartilhado entre todos os testes da sessão
        orders_store.save_orders(existing)


def test_list_audio_returns_json_structure(client):
    response = client.get("/list_audio")
    assert response.status_code == 200
    assert "audios" in response.get_json()


def test_download_audio_without_url_returns_400(logged_in_client):
    response = logged_in_client.post("/download_audio", json={})
    assert response.status_code == 400
    assert "error" in response.get_json()


def test_delete_audio_requires_login(client):
    response = client.delete("/delete_audio/nao-existe.mp3")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_delete_audio_returns_404_for_missing_file(logged_in_client):
    response = logged_in_client.delete("/delete_audio/nao-existe.mp3")
    assert response.status_code == 404


def test_visiting_lyrics_page_does_not_rewrite_unrelated_song_children(logged_in_client, app):
    # /lyrics/<id> chama save_songs() com a LISTA INTEIRA em toda visita,
    # só pra atualizar last_played de uma música — se o sync fosse um
    # "apaga tudo e reinsere" ingênuo, as seções desta música E das outras
    # 380 seriam recriadas (novos ids de linha) a cada visita.
    from models_songs import Song
    from songs_store import save_songs

    with app.app_context():
        save_songs([
            {"id": 90001, "title": "Música A", "author": "Autor", "sections": [{"title": "Verso", "text": ["la la"]}]},
            {"id": 90002, "title": "Música B", "author": "Autor", "sections": [{"title": "Verso", "text": ["la la"]}]},
        ])
        ids_before = {
            90001: [s.id for s in ucontrol_main.db.session.get(Song, 90001).sections],
            90002: [s.id for s in ucontrol_main.db.session.get(Song, 90002).sections],
        }

    logged_in_client.get("/lyrics/90001")
    logged_in_client.get("/lyrics/90001")

    with app.app_context():
        song_a = ucontrol_main.db.session.get(Song, 90001)
        last_played = song_a.last_played
        ids_after = {
            90001: [s.id for s in song_a.sections],
            90002: [s.id for s in ucontrol_main.db.session.get(Song, 90002).sections],
        }
        ucontrol_main.db.session.delete(song_a)
        ucontrol_main.db.session.delete(ucontrol_main.db.session.get(Song, 90002))
        ucontrol_main.db.session.commit()

    assert last_played is not None
    assert ids_after[90001] == ids_before[90001]
    assert ids_after[90002] == ids_before[90002]


def test_add_search_and_delete_song_round_trip(logged_in_client):
    add_response = logged_in_client.post(
        "/add_song",
        data={
            "title": "Música de Teste",
            "author": "Autor de Teste",
            "sections[]": ["Verso 1"],
            "texts[]": ["Letra de teste aqui"],
        },
    )
    assert add_response.status_code == 200

    search_response = logged_in_client.get("/search", query_string={"q": "Teste"})
    results = search_response.get_json()
    assert len(results) == 1
    song_id = results[0]["id"]
    assert results[0]["title"] == "Música de Teste"

    delete_response = logged_in_client.post(f"/delete_song/{song_id}")
    assert delete_response.status_code == 200


def test_view_song_falls_back_to_todays_order_when_no_order_found(app, client):
    # /view/<id> sem ?order= e sem nenhuma ordem referenciando a música
    # cai no fallback pra ordem de hoje (Fase 5.6 — antes lia
    # order_of_service.json). Cobertura nova: essa rota nunca teve teste.
    # Usa um item bíblico porque é o único pedaço de order_data que o
    # template view_song.html realmente renderiza (bible_items).
    import orders_store
    import songs_store
    from datetime import date

    today = date.today().isoformat()
    with app.app_context():
        existing_orders = orders_store.load_orders()
        existing_songs = songs_store.load_songs()
        song_id = songs_store.get_next_id(existing_songs)
        songs_store.save_songs(existing_songs + [{
            "id": song_id, "title": "Música Fallback", "author": "Autor",
            "sections": [{"title": "Verso", "text": ["la la"]}],
        }])
        orders_store.save_orders(dict(existing_orders, **{
            today: {"itens": [{"isBible": True, "reference": "Salmos 23:1", "title": "📖 Salmos 23:1"}]},
        }))

    response = client.get(f"/view/{song_id}")
    assert response.status_code == 200
    assert "Salmos 23:1".encode() in response.data

    with app.app_context():
        songs_store.save_songs(existing_songs)
        orders_store.save_orders(existing_orders)

    empty_search = client.get("/search", query_string={"q": "Teste"})
    assert empty_search.get_json() == []


def test_clock_and_help_pages_load(client):
    assert client.get("/clock").status_code == 200
    assert client.get("/help").status_code == 200


def test_historico_requires_login(client):
    response = client.get("/historico")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_api_historico_clear_empties_the_log(logged_in_client, app):
    from datetime import datetime

    from models import DisplayLog

    with app.app_context():
        ucontrol_main.db.session.add(DisplayLog(type="song", title="x", ts=datetime(2026, 1, 1)))
        ucontrol_main.db.session.commit()

    before = logged_in_client.get("/api/historico").get_json()
    assert len(before) >= 1

    clear_response = logged_in_client.delete("/api/historico")
    assert clear_response.get_json() == {"ok": True}

    after = logged_in_client.get("/api/historico").get_json()
    assert after == []


def test_append_display_log_inserts_a_real_row(app):
    from display_log import append_display_log, clear_display_log, load_display_log

    with app.app_context():
        clear_display_log()
        append_display_log({"type": "bible", "title": "João 3:16"})
        log = load_display_log()

    assert len(log) == 1
    assert log[0]["type"] == "bible"
    assert log[0]["title"] == "João 3:16"
    assert log[0]["url"] is None
    assert "ts" in log[0]


def test_load_display_log_returns_oldest_first(app):
    from display_log import append_display_log, clear_display_log, load_display_log

    with app.app_context():
        clear_display_log()
        append_display_log({"type": "image", "title": "primeira"})
        append_display_log({"type": "image", "title": "segunda"})
        log = load_display_log()

    assert [entry["title"] for entry in log] == ["primeira", "segunda"]


def test_settings_store_save_and_load_round_trip(app):
    from settings_store import load_settings, save_settings

    with app.app_context():
        save_settings({"bible_translation": "nvi"})
        save_settings({"display_bg_dark": True})
        settings = load_settings()

    assert settings["bible_translation"] == "nvi"
    assert settings["display_bg_dark"] is True


def test_image_path_and_saved_urls_round_trip(app):
    from blueprints.display import load_image_path, load_urls, save_image_path, save_urls

    with app.app_context():
        save_image_path("/static/imagens/foo.jpg")
        save_urls({"url1": "http://a", "url2": "http://b", "url3": "", "name1": "Câmera"})
        image = load_image_path()
        urls = load_urls()

    assert image == {"image": "/static/imagens/foo.jpg"}
    assert urls == {"url1": "http://a", "url2": "http://b", "url3": "",
                    "name1": "Câmera", "name2": "", "name3": ""}


def test_download_monitor_launcher_invalid_os_returns_400(client):
    response = client.get("/download-monitor-launcher/xyz")
    assert response.status_code == 400


def test_download_monitor_launcher_windows_returns_bat_attachment(client):
    response = client.get("/download-monitor-launcher/windows")
    assert response.status_code == 200
    assert ".bat" in response.headers["Content-Disposition"]


def test_connect_and_disconnect_update_client_tracking(app):
    import connections

    before = connections.socketio_clients_connected
    socket_client = _operator_socket(app)
    assert connections.socketio_clients_connected == before + 1

    socket_client.disconnect()
    assert connections.socketio_clients_connected == before


def test_identify_as_monitor_is_returned_by_get_monitor_sids(app):
    import connections

    socket_client = _operator_socket(app)
    socket_client.emit("identify", {"role": "monitor"})

    monitor_sids = connections._get_monitor_sids()
    assert len(monitor_sids) >= 1

    socket_client.disconnect()


def test_api_system_status_returns_expected_payload(logged_in_client):
    response = logged_in_client.get("/api/system/status")
    assert response.status_code == 200
    data = response.get_json()
    assert data["online"] is True
    assert data["port"] == 8585
    assert "socketio_clients" in data
    assert "socketio_by_role" in data


def test_orphaned_photo_navigation_handlers_were_removed():
    assert not hasattr(ucontrol_main, "handle_proxima_foto")
    assert not hasattr(ucontrol_main, "handle_anterior_foto")
    assert not hasattr(ucontrol_main, "obter_proxima_foto")


def test_list_image_folders_always_includes_root(client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    response = client.get("/api/image_folders")
    assert response.get_json() == ["Geral"]


def test_create_image_folder_and_see_it_listed(logged_in_client, tmp_path, monkeypatch):
    (tmp_path / "static" / "imagens").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)

    create_response = logged_in_client.post("/api/image_folders", json={"name": "Louvor"})
    assert create_response.get_json()["status"] == "success"

    folders = logged_in_client.get("/api/image_folders").get_json()
    assert "Louvor" in folders


def test_upload_list_and_delete_image_round_trip(logged_in_client, tmp_path, monkeypatch):
    import io
    from PIL import Image

    (tmp_path / "static" / "imagens").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)

    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), color="red").save(buffer, format="PNG")
    buffer.seek(0)

    upload_response = logged_in_client.post(
        "/upload",
        data={"images": (buffer, "teste.png")},
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 200
    saved = upload_response.get_json()["saved_files"]
    assert saved == ["teste.png"]

    images = logged_in_client.get("/imagens").get_json()
    assert images == ["/static/imagens/teste.png"]

    delete_response = logged_in_client.post("/delete_image", json={"image_name": "teste.png"})
    assert delete_response.get_json()["status"] == "success"
    assert logged_in_client.get("/imagens").get_json() == []


def test_delete_image_and_folder_reject_path_traversal(logged_in_client, tmp_path, monkeypatch):
    (tmp_path / "static" / "imagens").mkdir(parents=True)
    (tmp_path / "secret.txt").write_text("não deveria ser apagado")
    monkeypatch.chdir(tmp_path)

    delete_response = logged_in_client.post(
        "/delete_image", json={"image_name": "../secret.txt"}
    )
    assert delete_response.get_json()["status"] == "error"
    assert (tmp_path / "secret.txt").exists()

    folder_response = logged_in_client.delete("/api/image_folders/..")
    assert folder_response.status_code == 400
    assert (tmp_path / "static" / "imagens").exists()


def test_orders_store_save_and_load_round_trip(app):
    import orders_store

    with app.app_context():
        existing = orders_store.load_orders()
        # itens vazio é semanticamente "sem itens" (igual .pop('itens', None)
        # do sistema antigo) — por isso não entra na chave de volta, mesmo
        # padrão de chave-ausente-quando-vazio das Fases 5.2/5.4. Usa um
        # item de verdade aqui pra testar fidelidade do conteúdo, não essa
        # ambiguidade de lista vazia.
        sample = dict(existing, **{
            "2026-08-09-teste-store": {"tipo": "culto", "itens": [{"title": "Abertura", "url": "/order"}]},
        })
        assert orders_store.save_orders(sample) is True
        assert orders_store.load_orders().get("2026-08-09-teste-store") == sample["2026-08-09-teste-store"]
        # limpeza — banco de teste é compartilhado entre todos os testes da
        # sessão, remove só a entrada criada aqui, preserva o resto
        orders_store.save_orders(existing)
        assert "2026-08-09-teste-store" not in orders_store.load_orders()


def test_api_order_save_get_delete_round_trip(logged_in_client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    payload = {
        "tipo": "culto",
        "horario_culto": "19:00",
        "dirigente": "João",
        "pregador": "Pedro",
        "itens": [{"title": "Abertura"}],
    }

    save_response = logged_in_client.post("/api/orders/2026-08-09", json=payload)
    assert save_response.status_code == 200

    get_response = logged_in_client.get("/api/orders/2026-08-09")
    assert get_response.get_json()["dirigente"] == "João"

    delete_response = logged_in_client.delete("/api/orders/2026-08-09")
    assert delete_response.status_code == 200
    assert logged_in_client.get("/api/orders/2026-08-09").status_code == 404


def test_api_template_save_and_delete_round_trip(logged_in_client):
    save_response = logged_in_client.post(
        "/api/templates", json={"name": "Padrão", "items": ["Abertura", "Louvor"]}
    )
    assert save_response.get_json()["ok"] is True
    assert "Padrão" in logged_in_client.get("/api/templates").get_json()

    logged_in_client.delete("/api/templates/Padrão")
    assert "Padrão" not in logged_in_client.get("/api/templates").get_json()


def test_api_section_template_save_and_delete_round_trip(logged_in_client):
    save_response = logged_in_client.post(
        "/api/section-templates/louvor",
        json={"name": "Padrão", "fields": {"dirigente": "João"}},
    )
    assert save_response.get_json()["ok"] is True
    section = logged_in_client.get("/api/section-templates/louvor").get_json()
    assert section["Padrão"]["dirigente"] == "João"

    logged_in_client.delete("/api/section-templates/louvor/Padrão")
    section = logged_in_client.get("/api/section-templates/louvor").get_json()
    assert "Padrão" not in section


def test_legacy_save_and_load_compat_routes(app, logged_in_client):
    # Fase 5.6: /save e /load (sem ?date=) miram a ordem de HOJE na tabela
    # Order, não mais o antigo order_of_service.json (bucket sem data).
    import orders_store
    from datetime import date

    today = date.today().isoformat()
    with app.app_context():
        existing = orders_store.load_orders()

    save_response = logged_in_client.post("/save", json=[{"title": "Item legado", "url": "/order"}])
    assert save_response.status_code == 200

    load_response = logged_in_client.get("/load")
    assert load_response.get_json() == [{"title": "Item legado", "url": "/order"}]

    with app.app_context():
        assert orders_store.load_orders().get(today, {}).get("itens") == [
            {"title": "Item legado", "url": "/order"}
        ]
        # limpeza — banco de teste é compartilhado entre todos os testes da sessão
        orders_store.save_orders(existing)


def test_order_pages_load_with_login(logged_in_client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert logged_in_client.get("/order").status_code == 200
    assert logged_in_client.get("/order/new").status_code == 200
    assert logged_in_client.get("/order/2026-08-09").status_code == 200


def test_agenda_calendar_dashboard_and_edit_selector_load(logged_in_client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert logged_in_client.get("/agenda").status_code == 200
    assert logged_in_client.get("/agenda/dashboard").status_code == 200
    assert logged_in_client.get("/agenda/edit").status_code == 200


def test_agenda_detail_shows_real_order_data(logged_in_client, app):
    import orders_store

    with app.app_context():
        existing = orders_store.load_orders()
        orders_store.save_orders(dict(existing, **{"2026-08-09": {"dirigente": "João", "itens": []}}))

    response = logged_in_client.get("/agenda/2026-08-09")
    assert response.status_code == 200
    assert "João".encode() in response.data

    with app.app_context():
        orders_store.save_orders(existing)


def test_agenda_edit_renders_same_template_as_order_edit(logged_in_client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    agenda_response = logged_in_client.get("/agenda/2026-08-09/edit")
    order_response = logged_in_client.get("/order/2026-08-09")
    assert agenda_response.status_code == 200
    assert order_response.status_code == 200
    # Mesmo template (order_edit.html) por trás das duas rotas — o bloco
    # de dados que só esse template gera confirma isso.
    marker = b'id="order-edit-data"'
    assert marker in agenda_response.data
    assert marker in order_response.data


def test_save_route_and_duplicate_home_route_were_removed(client):
    assert not hasattr(ucontrol_main, "save_route")
    assert not hasattr(ucontrol_main, "home")
    # index() é quem responde em / (sempre venceu; home() nunca era alcançável)
    response = client.get("/")
    assert response.status_code == 302


def test_show_image_updates_display_state_module_live(app):
    # display_state._current_display_mode é reatribuído (não mutado em
    # lugar) pelos handlers — precisa ler via módulo pra pegar o valor
    # vivo, não um `from display_state import ...` que congelaria no
    # momento do import (mesma pegadinha do connections.py).
    import display_state

    socket_client = _operator_socket(app)
    socket_client.emit("show_image", {"image_url": "/static/imagens/teste.jpg"})

    assert display_state._current_display_mode == "image"
    assert display_state._last_image_url == "/static/imagens/teste.jpg"
    socket_client.disconnect()


def test_show_video_updates_display_state_module_live(app):
    import display_state

    socket_client = _operator_socket(app)
    socket_client.emit("show_video", {"video_url": "/static/videos/teste.mp4"})

    assert display_state._current_display_mode == "video"
    assert display_state._video_state["url"] == "/static/videos/teste.mp4"
    assert display_state._video_state["paused"] is False
    socket_client.disconnect()


def test_api_current_display_reflects_last_shown_image(app, client):
    socket_client = _operator_socket(app)
    socket_client.emit("show_image", {"image_url": "/static/imagens/atual.jpg"})
    socket_client.disconnect()

    response = client.get("/api/current-display")
    assert response.status_code == 200
    data = response.get_json()
    assert data["mode"] == "image"
    assert data["image_url"] == "/static/imagens/atual.jpg"


def test_api_monitor_sync_forwards_current_video_to_monitor_sids(app, logged_in_client):
    monitor_client = ucontrol_main.socketio.test_client(app)  # tela, sem login
    monitor_client.emit("identify", {"role": "monitor"})

    presenter_client = _operator_socket(app)
    presenter_client.emit("show_video", {"video_url": "/static/videos/sync.mp4"})
    presenter_client.disconnect()

    monitor_client.get_received()  # limpa os eventos já recebidos até aqui
    response = logged_in_client.post("/api/monitor/sync")
    assert response.status_code == 200

    received = monitor_client.get_received()
    events = [msg["name"] for msg in received]
    assert "display_monitor_video" in events
    payload = next(msg["args"][0] for msg in received if msg["name"] == "display_monitor_video")
    assert payload["video_url"] == "/static/videos/sync.mp4"

    monitor_client.disconnect()


def test_screen_share_offer_relays_to_target_sid(app):
    import connections

    # test_client.eio_sid é o sid de transporte do engine.io, não o
    # request.sid que os handlers usam — pega o sid de verdade via
    # connections.socketio_clients_by_sid (populado pelo handler connect),
    # conectando o receiver sozinho primeiro pra identificar sem ambiguidade.
    receiver = ucontrol_main.socketio.test_client(app)  # tela, sem login
    receiver_sid = next(iter(connections.socketio_clients_by_sid))
    sender = _operator_socket(app)

    sender.emit("screen_share_offer", {"sdp": "fake-sdp", "target_sid": receiver_sid})

    received = receiver.get_received()
    events = [msg["name"] for msg in received]
    assert "screen_share_offer" in events
    payload = next(msg["args"][0] for msg in received if msg["name"] == "screen_share_offer")
    assert payload["sdp"] == "fake-sdp"

    sender.disconnect()
    receiver.disconnect()


def test_update_lyrics_broadcasts_and_updates_display_state(app):
    import display_state

    socket_client = _operator_socket(app)
    socket_client.emit("update_lyrics", {"title": "Seção 1", "text": "Letra de teste"})

    received = socket_client.get_received()
    events = [msg["name"] for msg in received]
    assert "update_lyrics" in events
    assert display_state._current_display_mode == "lyrics"
    assert display_state._last_lyrics_payload == {"title": "Seção 1", "text": "Letra de teste"}
    socket_client.disconnect()


def test_api_lyrics_returns_song_data(logged_in_client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    add_response = logged_in_client.post(
        "/add_song",
        data={
            "title": "Música de Letras",
            "author": "Autor",
            "sections[]": ["Verso 1"],
            "texts[]": ["Letra aqui"],
        },
    )
    assert add_response.status_code == 200
    songs = logged_in_client.get("/search", query_string={"q": "Letras"}).get_json()
    song_id = songs[0]["id"]

    response = logged_in_client.get(f"/api/lyrics/{song_id}")
    assert response.status_code == 200
    assert response.get_json()["title"] == "Música de Letras"


# ── "Preparar vs. liberar" (StagedContent) ─────────────────────────────

def test_stage_content_persists_and_broadcasts(app):
    from extensions import db
    from models import StagedContent

    socket_client = _operator_socket(app)
    socket_client.emit("stage_content", {
        "kind": "image",
        "payload": {"image_url": "/static/imagens/preparada.jpg"},
        "label": "foto de teste",
    })

    received = socket_client.get_received()
    events = [msg["name"] for msg in received]
    assert "content_staged" in events

    with app.app_context():
        staged = db.session.get(StagedContent, "image")
        assert staged is not None
        assert staged.payload == {"image_url": "/static/imagens/preparada.jpg"}
        assert staged.label == "foto de teste"

    socket_client.disconnect()


def test_release_staged_content_dispatches_to_real_handler_and_clears_staging(app):
    import display_state
    from extensions import db
    from models import StagedContent

    socket_client = _operator_socket(app)
    socket_client.emit("stage_content", {
        "kind": "image",
        "payload": {"image_url": "/static/imagens/liberar.jpg"},
        "label": "foto",
    })
    socket_client.get_received()

    socket_client.emit("release_staged_content", {"kind": "image"})
    received = socket_client.get_received()
    events = [msg["name"] for msg in received]
    # "display_image" só existe se release chamou o handle_show_image de
    # verdade (não uma reimplementação) — é o mesmo evento que o teste de
    # show_image direto (test_show_image_updates_display_state_module_live)
    # verifica indiretamente via display_state.
    assert "content_released" in events
    assert display_state._last_image_url == "/static/imagens/liberar.jpg"
    assert display_state._current_display_mode == "image"

    with app.app_context():
        assert db.session.get(StagedContent, "image") is None

    socket_client.disconnect()


def test_discard_staged_content_clears_without_going_live(app):
    import display_state
    from extensions import db
    from models import StagedContent

    socket_client = _operator_socket(app)
    socket_client.emit("stage_content", {
        "kind": "lyrics",
        "payload": {"title": "Nunca vai ao ar", "text": "..."},
        "label": "letra",
    })
    socket_client.get_received()

    socket_client.emit("discard_staged_content", {"kind": "lyrics"})
    received = socket_client.get_received()
    events = [msg["name"] for msg in received]
    assert "content_discarded" in events
    assert "update_lyrics" not in events  # não foi ao ar
    assert display_state._current_display_mode != "lyrics"

    with app.app_context():
        assert db.session.get(StagedContent, "lyrics") is None

    socket_client.disconnect()


def test_staged_content_kinds_do_not_clobber_each_other(app):
    import display_state
    from extensions import db
    from models import StagedContent

    # Reseta o live-follow (ver blueprints/display.py, _release_kind) —
    # sem isto, um teste anterior que liberou "image" deixaria esse kind
    # marcado como já confirmado, e o stage_content("image") abaixo iria
    # direto ao ar em vez de só preparar, quebrando a premissa deste teste
    # (que é justamente checar que PREPARAR dois kinds ao mesmo tempo não
    # atropela um ao outro).
    display_state._live_follow_kind = None
    import telao
    telao.set_public('clock', '/clock')   # nem foto nem letra na tela
    telao.set_monitor(None)

    socket_client = _operator_socket(app)
    socket_client.emit("stage_content", {"kind": "image", "payload": {"image_url": "/foto.jpg"}, "label": "foto"})
    socket_client.emit("stage_content", {"kind": "lyrics", "payload": {"title": "Refrão"}, "label": "letra"})
    socket_client.get_received()

    with app.app_context():
        assert db.session.get(StagedContent, "image") is not None
        assert db.session.get(StagedContent, "lyrics") is not None

    socket_client.disconnect()


def test_api_staged_content_reflects_current_state(app, client):
    socket_client = _operator_socket(app)
    socket_client.emit("stage_content", {
        "kind": "bible",
        "payload": {"reference": "João 3:16", "text": "..."},
        "label": "João 3:16",
    })
    socket_client.disconnect()

    response = client.get("/api/staged-content")
    assert response.status_code == 200
    data = response.get_json()
    assert "bible" in data
    assert data["bible"]["label"] == "João 3:16"


def test_get_next_order_skips_empty_calendar_rows(app):
    # O calendário cria linhas vazias pras datas sem culto marcado — o card
    # "Próximo culto" do dashboard tem que pular essas.
    import orders_store
    from extensions import db
    from models_orders import Order

    with app.app_context():
        db.session.add(Order(date="2099-01-02"))
        db.session.commit()
        orders_store.save_order("2099-01-03", {"tipo": "Culto Dominical", "itens": [{"title": "Abertura", "url": "/order"}]})

        found_date, order = orders_store.get_next_order("2099-01-01")
        assert found_date == "2099-01-03"
        assert order["tipo"] == "Culto Dominical"
        assert orders_store.get_next_order("2099-01-04") == (None, {})

        orders_store.save_order("2099-01-03", {})
        db.session.delete(db.session.get(Order, "2099-01-02"))
        db.session.commit()


def test_api_culto_proximo_splits_song_artist_and_builds_team(app, logged_in_client, monkeypatch):
    import datetime as real_datetime
    import blueprints.orders as orders_module
    import orders_store

    class FakeDate(real_datetime.date):
        @classmethod
        def today(cls):
            return cls(2099, 2, 1)

    monkeypatch.setattr(orders_module, "date", FakeDate)
    with app.app_context():
        orders_store.save_order("2099-02-03", {
            "tipo": "Culto Dominical",
            "horario_culto": "18:40",
            "dirigente": "Miquéias",
            "itens": [
                {"title": "Ele É Exaltado | Adhemar de Campos", "url": "/lyrics/349"},
                {"isBible": True, "reference": "Salmos 23", "title": "Salmos 23"},
            ],
            "louvor": {"cantores": ["Bel"], "musicos": [{"instrumento": "Teclado", "pessoa": "Jefferson"}]},
            "midia": {"lider": "Vitor", "palco": "", "visual": "Thaynara"},
        })

    data = logged_in_client.get("/api/culto/proximo").get_json()
    assert data["status"] == "ok"
    assert data["date"] == "2099-02-03"
    assert data["days_until"] == 2
    assert data["horario"] == "18:40"
    assert data["itens"][0] == {"kind": "song", "title": "Ele É Exaltado", "artist": "Adhemar de Campos", "url": "/lyrics/349"}
    assert data["itens"][1]["kind"] == "bible"
    assert {"funcao": "Mídia", "pessoa": "Vitor"} in data["equipe"]
    assert {"funcao": "Teclado", "pessoa": "Jefferson"} in data["equipe"]
    assert {"funcao": "Vozes", "pessoa": "Bel"} in data["equipe"]
    assert not any(e["funcao"] == "Palco" for e in data["equipe"])  # vazio não entra

    with app.app_context():
        orders_store.save_order("2099-02-03", {})


def test_group_activity_merges_consecutive_repeats():
    from blueprints.auth import group_activity

    log = [
        {"type": "image", "title": "/static/imagens/Geral/boas%20vindas.jpg", "ts": "2026-10-08T18:00:00"},
        {"type": "song", "title": "Ele É Exaltado", "ts": "2026-10-08T19:08:00"},
        {"type": "song", "title": "Ele É Exaltado", "ts": "2026-10-08T19:09:00"},
        {"type": "song", "title": "Ele É Exaltado", "ts": "2026-10-08T19:11:00"},
        {"type": "image", "title": "/lyrics/display + /clock", "ts": "2026-10-08T19:12:00"},
    ]
    groups = group_activity(log)

    assert [g["count"] for g in groups] == [1, 3, 1]
    assert groups[0]["target"] == "Cifra + Relógio"
    assert groups[1]["action"] == "Exibiu a letra de"
    assert (groups[1]["first_ts"], groups[1]["last_ts"]) == ("2026-10-08T19:08:00", "2026-10-08T19:11:00")
    assert groups[2]["target"] == "boas vindas.jpg"


def test_redesigned_pages_render_with_theme_and_nav(logged_in_client):
    for path in ("/dashboard", "/control", "/usuarios", "/songs", "/new_song", "/sheet_editor",
                 "/order", "/order/new", "/historico", "/links", "/help", "/bible",
                 "/bible/jo", "/bible/jo/3", "/bible/control/jo/3", "/presenter",
                 "/agenda/dashboard", "/agenda/edit", "/agenda/2099-01-01/admin",
                 "/agenda/2099-01-01/louvor"):
        response = logged_in_client.get(path)
        assert response.status_code == 200, path
        html = response.get_data(as_text=True)
        assert 'class="uc-app"' in html, path
        assert "css/app.css" in html, path
        assert 'class="site-nav"' in html, path
        assert "site-nav-link live" in html, path  # botão do Controle


def test_projector_screens_stay_outside_theme_and_nav(client):
    # Telas de projetor (monitor, cronômetro, relógio) têm visual próprio
    # e ninguém clica nelas — não recebem o app.css nem a barra.
    for path in ("/monitor", "/countdown", "/clock"):
        html = client.get(path).get_data(as_text=True)
        assert "css/app.css" not in html, path
        assert 'class="site-nav"' not in html, path


def test_public_pages_use_theme_without_operator_menu(client):
    # Páginas que a equipe abre sem login (agenda, cifras): tema novo, mas a
    # barra só mostra a marca e "Entrar" — nada do menu do operador.
    for path in ("/login", "/agenda", "/agenda/2099-01-01", "/view"):
        response = client.get(path)
        assert response.status_code == 200, path
        html = response.get_data(as_text=True)
        assert 'class="uc-app"' in html, path
        assert "site-nav-link live" not in html, path



def test_anonymous_socket_cannot_control_screens(app):
    """Sem login, comandos de controle do telão são ignorados — antes
    qualquer visitante do site podia trocar o que aparece no projetor."""
    screen = ucontrol_main.socketio.test_client(app)
    screen.emit("identify", {"role": "public"})
    screen.get_received()

    intruder = ucontrol_main.socketio.test_client(app)
    for event, data in (("show_image", {"image_url": "https://exemplo.com"}),
                        ("set_ticker", {"text": "invasão", "target": "all"}),
                        ("blackout", {"active": True}),
                        ("show_two_urls", {"url1": "/clock", "url2": "/img"}),
                        ("start_countdown", {"duration_seconds": 60}),
                        ("toggle_video_pause", {"paused": False})):
        intruder.emit(event, data)

    events = [msg["name"] for msg in screen.get_received()]
    for blocked in ("display_image", "ticker_update", "blackout", "display_two_urls",
                    "countdown_start", "video_pause"):
        assert blocked not in events, blocked

    intruder.disconnect()
    screen.disconnect()


def test_screens_without_login_can_still_sync_video_reset(app):
    """/monitor e /img não têm login e pedem "pausado no início" ao
    carregar um vídeo — essa forma específica continua aceita."""
    screen = ucontrol_main.socketio.test_client(app)
    screen.emit("identify", {"role": "monitor"})
    screen.emit("toggle_video_pause", {"paused": True})
    events = [msg["name"] for msg in screen.get_received()]
    assert "video_pause" in events
    screen.disconnect()


def test_logged_operator_socket_controls_screens(app):
    screen = ucontrol_main.socketio.test_client(app)
    operator = _operator_socket(app)
    operator.emit("blackout", {"active": False})
    events = [msg["name"] for msg in screen.get_received()]
    assert "blackout" in events
    operator.disconnect()
    screen.disconnect()


def test_presenter_requires_login(client):
    response = client.get("/presenter")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_blackout_state_is_sent_to_screen_that_connects_later(app):
    operator = _operator_socket(app)
    operator.emit("blackout", {"active": True})

    late_screen = ucontrol_main.socketio.test_client(app)
    late_screen.emit("identify", {"role": "public"})
    received = late_screen.get_received()
    assert any(m["name"] == "blackout" and m["args"][0]["active"] for m in received)

    operator.emit("blackout", {"active": False})
    other = ucontrol_main.socketio.test_client(app)
    other.emit("identify", {"role": "public"})
    assert not any(m["name"] == "blackout" for m in other.get_received())

    for c in (operator, late_screen, other):
        c.disconnect()


def test_display_public_has_full_screen_blackout_layer(client):
    html = client.get("/lyrics/display/public").get_data(as_text=True)
    assert 'id="blackout-layer"' in html


def _release_photo(op, url):
    op.emit("stage_content", {"kind": "image", "payload": {"image_url": url}, "label": "foto"})
    op.emit("release_staged_content", {"kind": "image"})
    op.get_received()


def test_control_switching_to_other_source_ends_photo_live_follow(app):
    """Depois de liberar uma foto do /presenter, pôr o relógio no telão
    pelo /control faz a próxima foto voltar a esperar o Liberar — antes
    ela entrava por cima do relógio na hora."""
    import display_state

    op = _operator_socket(app)
    _release_photo(op, "/static/imagens/a.jpg")
    assert display_state._live_follow_kind == "image"

    op.emit("show_image", {"image_url": "/clock"})
    assert display_state._live_follow_kind is None

    op.emit("stage_content", {"kind": "image", "payload": {"image_url": "/static/imagens/b.jpg"}, "label": "b"})
    events = [m["name"] for m in op.get_received()]
    assert "content_staged" in events
    assert display_state._last_image_url == "/clock"
    op.emit("discard_staged_content", {"kind": "image"})
    op.disconnect()


def test_control_staying_on_photos_keeps_live_follow(app):
    import display_state

    op = _operator_socket(app)
    _release_photo(op, "/static/imagens/a.jpg")
    op.emit("show_image", {"image_url": "/static/imagens/galeria.jpg"})
    op.emit("show_image", {"image_url": "/img"})
    assert display_state._live_follow_kind == "image"

    op.emit("stage_content", {"kind": "image", "payload": {"image_url": "/static/imagens/c.jpg"}, "label": "c"})
    assert display_state._last_image_url == "/static/imagens/c.jpg"

    op.emit("show_two_urls", {"url1": "/lyrics/display", "url2": "/clock"})
    assert display_state._live_follow_kind is None
    op.disconnect()


def _telao_status_events(client):
    return [m["args"][0] for m in client.get_received() if m["name"] == "telao_status"]


def test_telao_status_describes_live_staged_and_follow(app):
    import display_state

    op = _operator_socket(app)
    op.emit("show_image", {"image_url": "/clock"})
    status = _telao_status_events(op)[-1]
    assert status["live"]["label"] == "Relógio"
    assert status["monitor"] is None

    op.emit("stage_content", {"kind": "image", "payload": {"image_url": "/static/imagens/Culto%20Jovem.jpg"},
                              "label": "Culto Jovem.jpg"})
    status = _telao_status_events(op)[-1]
    assert {"kind": "image", "label": "Culto Jovem.jpg"} in status["staged"]

    op.emit("release_staged_content", {"kind": "image"})
    status = _telao_status_events(op)[-1]
    assert status["live"]["label"] == "Foto · Culto Jovem.jpg"
    assert not any(item["kind"] == "image" for item in status["staged"])
    assert status["follow"] == "image" and status["follow_label"] == "Fotos"

    op.emit("stop_live_follow", {})
    status = _telao_status_events(op)[-1]
    assert status["follow"] is None
    assert display_state._live_follow_kind is None
    op.disconnect()


def test_telao_status_shows_monitor_when_split_or_public_only(app):
    op = _operator_socket(app)
    op.emit("show_image", {"image_url": "/clock"})
    op.emit("show_two_urls", {"url1": "/lyrics/display", "url2": "/clock"})
    status = _telao_status_events(op)[-1]
    assert status["live"]["label"] == "Relógio"  # a tela pública não divide
    assert status["monitor"]["label"] == "Tela dividida · Cifra + Relógio"

    op.emit("show_image", {"image_url": "/bible/display"})
    assert _telao_status_events(op)[-1]["monitor"] is None

    op.emit("show_image", {"image_url": "/clock", "exclude_monitor": True})
    status = _telao_status_events(op)[-1]
    assert status["live"]["label"] == "Relógio"
    assert status["monitor"]["label"].startswith("Bíblia")
    op.disconnect()


def test_telao_status_route_requires_login(client, logged_in_client):
    assert client.get("/api/telao-status").status_code == 302
    data = logged_in_client.get("/api/telao-status").get_json()
    assert set(data) >= {"live", "monitor", "staged", "mode", "follow", "blackout"}


def test_telao_bar_only_for_logged_operator(client, logged_in_client):
    assert 'id="telao-bar"' in logged_in_client.get("/dashboard").get_data(as_text=True)
    assert 'id="telao-bar"' not in client.get("/agenda").get_data(as_text=True)
    assert 'id="telao-bar"' not in client.get("/lyrics/display/public").get_data(as_text=True)


def test_anonymous_cannot_stop_live_follow(app):
    import display_state

    op = _operator_socket(app)
    op.emit("stage_content", {"kind": "image", "payload": {"image_url": "/static/imagens/x.jpg"}, "label": "x"})
    op.emit("release_staged_content", {"kind": "image"})
    assert display_state._live_follow_kind == "image"
    intruder = ucontrol_main.socketio.test_client(app)
    intruder.emit("stop_live_follow", {})
    assert display_state._live_follow_kind == "image"
    op.emit("stop_live_follow", {})
    intruder.disconnect()
    op.disconnect()


def test_monitor_sync_copies_lyrics_released_from_other_page(app, client, logged_in_client):
    """Letra liberada pela página da música em Modo Público: o Sincronizar
    leva a letra pro monitor (antes mandava a última fonte do /control)."""
    monitor = ucontrol_main.socketio.test_client(app)  # tela, sem login
    monitor.emit("identify", {"role": "monitor"})
    op = _operator_socket(app)
    op.emit("show_image", {"image_url": "/bible/display"})
    op.emit("set_display_mode", {"exclude_monitor": True, "monitor_only": False})
    op.emit("stage_content", {"kind": "lyrics", "label": "Teu Povo — Coro",
                              "payload": {"songId": 1, "sectionIndex": 0, "title": "Coro", "text": "x", "type": "text"}})
    op.emit("release_staged_content", {"kind": "lyrics"})
    status = [m["args"][0] for m in op.get_received() if m["name"] == "telao_status"][-1]
    assert status["live"]["kind"] == "lyrics"
    assert status["monitor"]["label"].startswith("Bíblia")  # monitor congelado

    monitor.get_received()
    assert client.post("/api/monitor/sync").status_code in (302, 401)  # exige login
    assert logged_in_client.post("/api/monitor/sync").status_code == 200
    sent = [m["args"][0]["url"] for m in monitor.get_received() if m["name"] == "display_monitor_content"]
    assert sent == ["/lyrics/display/public"]

    op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
    op.emit("stop_live_follow", {})
    op.disconnect()
    monitor.disconnect()


def test_bible_display_can_fetch_current_verse(app, client):
    op = _operator_socket(app)
    op.emit("update_bible", {"reference": "João 3:16", "text": "16 Porque Deus amou o mundo",
                             "verses": [{"num": 16, "text": "Porque Deus amou o mundo"}]})
    data = client.get("/api/bible/current").get_json()  # tela sem login
    assert data["reference"] == "João 3:16"
    op.disconnect()


def test_telao_ping_acks_for_operator(app):
    op = _operator_socket(app)
    assert op.emit("telao_ping", {}, callback=True) == {"ok": True}
    op.disconnect()


def test_command_pages_mark_their_command_area(logged_in_client):
    for path in ("/control", "/presenter", "/lyrics/1", "/bible/jo/3", "/bible/control/jo/3/16-16"):
        html = logged_in_client.get(path).get_data(as_text=True)
        assert "data-telao-commands" in html, path


def test_telao_status_shows_slideshow_ticker_countdown_and_stops_them(app):
    op = _operator_socket(app)
    control = _operator_socket(app)
    op.emit("stop_ticker", {"target": "all"})  # outro teste pode ter deixado aviso ligado
    control.emit("slideshow_state", {"running": True, "interval": 7, "count": 12})
    op.emit("set_ticker", {"text": "Bem-vindos", "target": "public"})
    op.emit("start_countdown", {"duration_seconds": 300, "label": "Culto"})
    status = _telao_status_events(op)[-1]
    assert status["slideshow"] == {"interval": 7, "count": 12}
    assert status["ticker"] == {"text": "Bem-vindos", "targets": ["public"]}
    assert status["countdown"]["label"] == "Culto" and not status["countdown"]["paused"]

    control.get_received()
    op.emit("stop_slideshow", {})
    assert any(m["name"] == "slideshow_stop" for m in control.get_received())
    op.emit("stop_ticker", {"target": "all"})
    op.emit("reset_countdown", {})
    status = _telao_status_events(op)[-1]
    assert status["slideshow"] is None and status["ticker"] is None and status["countdown"] is None
    op.disconnect()
    control.disconnect()


def test_slideshow_forgotten_when_its_control_tab_closes(app):
    op = _operator_socket(app)
    control = _operator_socket(app)
    control.emit("slideshow_state", {"running": True, "interval": 5, "count": 3})
    control.disconnect()
    op.emit("telao_ping", {})
    import telao
    assert telao.build_status()["slideshow"] is None
    op.disconnect()


def test_panel_song_api_and_order_bible_urls(logged_in_client):
    data = logged_in_client.get("/api/painel/musica").get_json()
    assert set(data) == {"song_id", "order_date"}


def test_embedded_pages_hide_nav_via_client_detection(logged_in_client):
    html = logged_in_client.get("/presenter").get_data(as_text=True)
    assert "uc-embed" in html  # script que detecta iframe e esconde menu/barra


def test_control_has_panel_tools_and_mode_choice(logged_in_client):
    html = logged_in_client.get("/control").get_data(as_text=True)
    for marker in ('id="ctl-panel"', 'data-page="musica"', 'data-page="biblia"', 'data-page="fotos"',
                   'data-page="ordem"', 'data-open-tool="timer"', 'data-open-tool="split"',
                   'data-mode="public"', 'data-mode="monitor"', 'id="name1"', 'data-slot="left"'):
        assert marker in html, marker
    assert "monitor-mode-btn" not in html  # interruptores antigos saíram


def test_save_urls_requires_login_and_keeps_names(app, client, logged_in_client):
    assert client.post("/save_urls", json={"url1": "x"}).status_code in (302, 400, 401)
    r = logged_in_client.post("/save_urls", json={"url1": "http://cam", "name1": "Câmera",
                                                  "url2": "", "url3": ""})
    assert r.get_json()["status"] == "success"
    assert logged_in_client.get("/load_urls").get_json()["name1"] == "Câmera"


def test_single_gallery_lives_in_presenter(logged_in_client):
    control = logged_in_client.get("/control").get_data(as_text=True)
    assert 'id="imageGallery"' not in control  # galeria do Controle saiu
    assert 'data-open-page="fotos"' in control
    presenter = logged_in_client.get("/presenter").get_data(as_text=True)
    for marker in ('id="imageGallery"', 'id="btnSlideshow"', 'id="slideshowInterval"', 'id="newFolderName"',
                   'id="deleteFolderBtn"', 'id="toggleSelectBtn"', 'id="deleteSelectedBtn"', 'id="fileInput"'):
        assert marker in presenter, marker


def test_each_machine_has_its_own_display_mode(app):
    """Dois /control em PCs diferentes (cookie uc_device diferente): um em
    Só público, outro em Só monitor. Trocar o modo num não troca no outro,
    e cada um manda pra sua tela."""
    import display_state
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    pc_publico = _operator_socket(app)
    pc_palco = _operator_socket(app)
    pc_publico.emit("set_display_mode", {"exclude_monitor": True, "monitor_only": False})
    pc_palco.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": True})
    status = [m["args"][0] for m in pc_publico.get_received() if m["name"] == "telao_status"][-1]
    assert sorted(status["modes"].values()) == ["monitor", "public"]

    monitor.get_received()
    pc_publico.emit("show_image", {"image_url": "/static/imagens/publico.jpg", "use_active_mode": True})
    assert display_state._last_image_url == "/static/imagens/publico.jpg"
    assert not [m for m in monitor.get_received() if m["name"] == "display_monitor_content"]

    pc_palco.emit("show_image", {"image_url": "/static/imagens/palco.jpg", "use_active_mode": True})
    got = [m["args"][0]["url"] for m in monitor.get_received() if m["name"] == "display_monitor_content"]
    assert got == ["/static/imagens/palco.jpg"]
    assert display_state._last_image_url == "/static/imagens/publico.jpg"  # público não mudou

    for op in (pc_publico, pc_palco):
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.disconnect()
    assert display_state._device_modes == {}
    monitor.disconnect()


def test_slideshow_show_image_respects_server_mode(app):
    import display_state
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    op = _operator_socket(app)
    op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": True})
    before = display_state._last_image_url
    monitor.get_received()
    op.emit("show_image", {"image_url": "/static/imagens/slide.jpg", "use_active_mode": True})
    got = [m for m in monitor.get_received() if m["name"] == "display_monitor_content"]
    assert got and got[0]["args"][0]["url"] == "/static/imagens/slide.jpg"
    assert display_state._last_image_url == before  # tela pública não mudou
    op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
    op.disconnect()
    monitor.disconnect()


def test_letra_route_gives_current_lyrics_to_screens_opening_later(app, client):
    op = _operator_socket(app)
    op.emit("update_lyrics", {"title": "Coro", "text": "Santo, santo", "songId": 1, "type": "text"})
    op.emit("show_image", {"image_url": "/lyrics/display/public"})
    data = client.get("/api/current-display").get_json()
    assert data["mode"] == "lyrics" and data["text"] == "Santo, santo"
    op.disconnect()


def test_cifra_shows_lyrics_to_public_and_cifra_on_monitor(app):
    op = _operator_socket(app)
    op.emit("show_image", {"image_url": "/lyrics/display"})
    status = _telao_status_events(op)[-1]
    assert status["live"]["kind"] == "lyrics"
    assert status["live"]["url"] == "/lyrics/display/public"
    assert status["monitor"]["kind"] == "cifra"
    op.emit("show_image", {"image_url": "/lyrics/display/public"})
    assert _telao_status_events(op)[-1]["monitor"] is None
    op.disconnect()


def test_telao_status_gives_monitor_urls_for_preview(app):
    op = _operator_socket(app)
    op.emit("show_image", {"image_url": "/clock"})
    op.emit("show_monitor_content", {"url": "/lyrics/display"})
    mon = _telao_status_events(op)[-1]["monitor"]
    assert mon["kind"] == "cifra" and mon["url"] == "/lyrics/display"
    op.emit("show_two_urls", {"url1": "/lyrics/display", "url2": "/clock"})
    mon = _telao_status_events(op)[-1]["monitor"]
    assert (mon["url1"], mon["url2"]) == ("/lyrics/display", "/clock")
    op.disconnect()


def test_screens_load_shared_zoom_fit(client):
    for path in ("/monitor", "/lyrics/display/public", "/img"):
        assert "js/zoom_fit.js" in client.get(path).get_data(as_text=True), path


def test_agenda_edit_has_whatsapp_share_buttons(logged_in_client):
    html = logged_in_client.get("/agenda/edit").get_data(as_text=True)
    for marker in ('id="copyWhatsBtn"', 'id="shareMenu"', 'data-format="grid"', 'data-format="light"',
                   'data-format="dark"', 'id="saveImgBtn"', "js/pages/agenda_share.js"):
        assert marker in html, marker


def test_culto_keys_and_current_culto_of_the_day(app):
    from datetime import datetime

    import orders_store as s

    assert (s.key_date("2026-10-04-2"), s.key_seq("2026-10-04-2")) == ("2026-10-04", 2)
    assert (s.key_seq("2026-10-04"), s.make_key("2026-10-04", 1), s.make_key("2026-10-04", 3)) == (1, "2026-10-04", "2026-10-04-3")
    with app.app_context():
        assert s.next_free_key("2031-05-04") == "2031-05-04"
        s.save_order("2031-05-04", {"tipo": "Culto Dominical", "horario_culto": "18:40",
                                    "eventos": [{"titulo": "Ensaio", "hora": "16:00"}]})
        assert s.next_free_key("2031-05-04") == "2031-05-04-2"
        s.save_order("2031-05-04-2", {"tipo": "Escola Bíblica", "horario_culto": "09:00"})
        try:
            # Em ordem de horário, não de chave
            assert [k for k, _ in s.day_orders("2031-05-04")] == ["2031-05-04-2", "2031-05-04"]
            assert s.next_free_key("2031-05-04") == "2031-05-04-3"
            # De manhã é o da manhã; a partir de 1h antes do da noite, o da noite
            assert s.current_key("2031-05-04", datetime(2031, 5, 4, 10, 0)) == "2031-05-04-2"
            assert s.current_key("2031-05-04", datetime(2031, 5, 4, 17, 50)) == "2031-05-04"
            assert s.current_key("2031-05-04", datetime(2031, 5, 4, 7, 0)) == "2031-05-04-2"
            # Outro dia: o primeiro pelo horário
            assert s.current_key("2031-05-04", datetime(2031, 5, 1, 20, 0)) == "2031-05-04-2"
            assert s.get_next_order("2031-05-04", datetime(2031, 5, 1, 20, 0))[0] == "2031-05-04-2"
            assert s.current_key("2031-05-05") == "2031-05-05"
        finally:
            s.save_order("2031-05-04", {})
            s.save_order("2031-05-04-2", {})


def test_second_culto_has_own_pages_and_no_events(logged_in_client):
    c = logged_in_client
    c.post("/api/orders/2031-06-01", json={"tipo": "Culto Dominical", "horario_culto": "18:40",
                                           "eventos": [{"titulo": "Ensaio do louvor", "hora": "16:00"}]})
    try:
        # "Novo culto" abre o louvor da próxima chave livre
        r = c.get("/agenda/2031-06-01/novo-culto")
        assert r.status_code == 302 and r.headers["Location"].endswith("/agenda/2031-06-01-2/louvor")
        # Eventos são do dia: ignorados no 2º culto
        c.post("/api/orders/2031-06-01-2", json={"tipo": "Escola Bíblica", "horario_culto": "09:00",
                                                 "eventos": [{"titulo": "Não deveria ficar"}]})
        assert "eventos" not in c.get("/api/orders/2031-06-01-2").get_json()

        day = c.get("/agenda/2031-06-01/admin").get_data(as_text=True)
        assert day.index("Escola Bíblica") < day.index("Culto Dominical")
        assert "/agenda/2031-06-01-2/louvor" in day and "Ensaio do louvor" in day
        # Chave do 2º culto na URL da página do dia mostra o dia inteiro
        assert "Ensaio do louvor" in c.get("/agenda/2031-06-01-2").get_data(as_text=True)

        louvor = c.get("/agenda/2031-06-01-2/louvor").get_data(as_text=True)
        assert 'href="/agenda/2031-06-01/louvor"' in louvor and 'id="deleteCultoBtn"' in louvor
        assert c.get("/agenda/2031-06-01-2/eventos").headers["Location"].endswith("/agenda/2031-06-01/eventos")

        # Excluir o 1º culto mantém os eventos do dia
        assert c.delete("/api/cultos/2031-06-01").status_code == 200
        assert c.get("/api/orders/2031-06-01").get_json() == {
            "eventos": [{"titulo": "Ensaio do louvor", "hora": "16:00", "responsavel": "", "obs": ""}]}
        assert c.delete("/api/cultos/2031-06-01-2").status_code == 200
        assert c.get("/api/orders/2031-06-01-2").status_code == 404
    finally:
        c.delete("/api/orders/2031-06-01")
        c.delete("/api/orders/2031-06-01-2")


def test_agenda_calendars_load_shared_culto_helper(logged_in_client):
    for path in ("/agenda", "/agenda/edit"):
        assert "js/agenda_cultos.js" in logged_in_client.get(path).get_data(as_text=True), path


def test_tipo_culto_outro_keeps_custom_name(logged_in_client):
    import re

    c = logged_in_client
    c.post("/api/orders/2031-07-05", json={"tipo": "Culto de Missões", "horario_culto": "19:00"})
    try:
        for path in ("/order/2031-07-05", "/agenda/2031-07-05/louvor"):
            html = c.get(path).get_data(as_text=True)
            assert '<option value="Outro" selected>' in html, path
            campo = re.search(r'<input[^>]*id="(?:orderType|tipo)Outro"[^>]*>', html).group(0)
            assert 'value="Culto de Missões"' in campo and " hidden" not in campo, path
        # Tipo da lista: "Outro" fechado
        c.post("/api/orders/2031-07-05", json={"tipo": "Culto de Oração", "horario_culto": "19:30"})
        html = c.get("/order/2031-07-05").get_data(as_text=True)
        assert '<option value="Culto de Oração" selected>' in html and 'id="orderTypeOutro"' in html
    finally:
        c.delete("/api/orders/2031-07-05")


def test_view_order_marks_items_done(app, client, logged_in_client):
    import orders_store

    itens = [{"anyText": "Oração"}, {"title": "Hino | Autor", "url": "/lyrics/1"},
             {"isBible": True, "reference": "João 3:16", "title": "📖 João 3:16"}, {"anyText": "Oração"}]
    assert orders_store.item_ids(itens) == ["texto:Oração#1", "url:/lyrics/1#1", "biblia:João 3:16#1", "texto:Oração#2"]
    logged_in_client.post("/api/orders/2031-08-02", json={"tipo": "Culto Dominical", "itens": itens})
    try:
        # Visitante não marca; logado marca, todos veem
        assert client.post("/api/orders/2031-08-02/feitos", json={"item": "texto:Oração#2", "feito": True}).status_code in (302, 401)
        r = logged_in_client.post("/api/orders/2031-08-02/feitos", json={"item": "texto:Oração#2", "feito": True})
        assert r.get_json()["feitos"] == ["texto:Oração#2"]
        assert client.get("/api/orders/2031-08-02/feitos").get_json()["feitos"] == ["texto:Oração#2"]
        html = client.get("/view/order/2031-08-02").get_data(as_text=True)
        assert html.count(" feito\"") == 1 and 'data-item="texto:Oração#2"' in html
        assert 'button type="button" class="item-check"' not in html
        assert 'class="item-check" aria-pressed' in logged_in_client.get("/view/order/2031-08-02").get_data(as_text=True)
        # Editar a ordem não apaga a marcação
        logged_in_client.post("/api/orders/2031-08-02", json={"tipo": "Culto Dominical", "itens": itens[::-1]})
        assert client.get("/api/orders/2031-08-02/feitos").get_json()["feitos"] == ["texto:Oração#2"]
        logged_in_client.post("/api/orders/2031-08-02/feitos", json={"limpar": True})
        assert client.get("/api/orders/2031-08-02/feitos").get_json()["feitos"] == []
    finally:
        logged_in_client.delete("/api/orders/2031-08-02")


def test_items_going_on_screen_are_marked_done(app):
    from datetime import date

    import display_state
    import orders_store
    from blueprints import display

    today = date.today().isoformat()
    itens = [{"title": "Hino | A", "url": "/lyrics/7"},
             {"isBible": True, "reference": "Romanos 10:14-17", "title": "📖 Romanos 10:14-17"},
             {"anyText": "Oração"}, {"title": "Hino | A", "url": "/lyrics/7"}]
    with app.app_context():
        antes = orders_store.get_order(today)
        orders_store.save_order(today, {"tipo": "Culto Dominical", "itens": itens})
    display_state._live_follow_kind = None
    display._ultimo_feito_auto = None
    sock = _operator_socket(app)

    def no_ar(kind, payload):
        sock.emit("stage_content", {"kind": kind, "payload": payload, "label": "x"})
        sock.emit("release_staged_content", {"kind": kind})
        sock.get_received()
        with app.app_context():
            return orders_store.get_feitos(today)

    try:
        assert no_ar("lyrics", {"songId": 7, "sectionIndex": 0, "title": "V1"}) == ["url:/lyrics/7#1"]
        # Outra seção da mesma música não marca a 2ª vez que ela aparece na ordem
        assert no_ar("lyrics", {"songId": 7, "sectionIndex": 1, "title": "Refrão"}) == ["url:/lyrics/7#1"]
        # Versículo dentro da passagem da ordem marca a leitura
        assert no_ar("bible", {"reference": "Romanos 10:15", "text": "..."})[-1] == "biblia:Romanos 10:14-17#1"
        # A música volta depois de outra coisa: agora é a 2ª ocorrência
        assert no_ar("lyrics", {"songId": 7, "sectionIndex": 0, "title": "V1"})[-1] == "url:/lyrics/7#2"
        # Música fora da ordem não marca nada
        assert len(no_ar("lyrics", {"songId": 999, "sectionIndex": 0, "title": "X"})) == 3
    finally:
        sock.disconnect()
        display_state._live_follow_kind = None
        with app.app_context():
            orders_store.save_order(today, {})
            if antes:
                orders_store.save_order(today, antes)


def test_launcher_opens_chosen_screen_and_page(client, logged_in_client):
    padrao = client.get("/download-monitor-launcher/windows").get_data(as_text=True)
    assert "/img" in padrao and "--window-position" not in padrao and "UControl-Monitor" in padrao
    r = client.get("/download-monitor-launcher/windows?tela=publico&x=1930&y=10")
    corpo = r.get_data(as_text=True)
    assert "Tela-Publica" in r.headers["Content-Disposition"]
    assert "/lyrics/display/public" in corpo and "--window-position=1930,10 --start-fullscreen" in corpo
    assert "--window-position=-190,60" in client.get("/download-monitor-launcher/linux?tela=palco&x=-190&y=60").get_data(as_text=True)
    html = logged_in_client.get("/control").get_data(as_text=True)
    assert 'id="modal-datashow"' in html and "js/screen_picker.js" in html
    assert "js/tela_cheia.js" in client.get("/lyrics/display/public").get_data(as_text=True)


def test_every_bible_book_opens_in_every_translation(logged_in_client):
    # ARC/KJA e as estrangeiras usam abreviações em inglês ('jm' = Tiago,
    # 'act' = Atos): antes 17 livros davam "Livro não encontrado" nelas
    from blueprints.bible import load_bible_metadata

    meta = load_bible_metadata()
    falhas = []
    for t in meta["translations"]:
        for book in meta["books"]:
            for path in (f"/bible/{book['abbrev']}/1?translation={t}", f"/bible/control/{book['abbrev']}/1?translation={t}"):
                if logged_in_client.get(path).status_code != 200:
                    falhas.append(path)
    assert falhas == []
    # en_kjv: abreviações em inglês, como ARC/KJA (essas só existem na pasta
    # privada de quem tem o arquivo, ver paths.BIBLES_PRIVATE_DIR)
    assert logged_in_client.get("/bible/control/tg/1?translation=en_kjv").status_code == 200


def test_flash_shows_on_the_page_it_redirects_to(logged_in_client):
    c = logged_in_client
    for _ in range(3):
        r = c.get("/bible/nao-existe/1")
        assert r.status_code == 302
    html = c.get(r.headers["Location"]).get_data(as_text=True)
    # Os 3 avisos viram um toast só na página da Bíblia...
    assert html.count("encontrado!") == 1 and "ucToast" in html  # tojson escreve "n\\u00e3o"
    # ...e não sobram pra aparecer depois em /usuarios
    assert "encontrado!" not in c.get("/usuarios").get_data(as_text=True)


def test_bible_chapter_page_has_arrow_navigation_data(logged_in_client):
    import json
    import re

    html = logged_in_client.get("/bible/tg/1?translation=nbv").get_data(as_text=True)
    dados = json.loads(re.search(r'id="bible-chapter-data" type="application/json">(.*?)</script>', html).group(1))
    assert (dados["prevChapter"], dados["nextChapter"], dados["translation"]) == (None, 2, "nbv")
    assert "setas" in html


def test_hide_button_reaches_bible_screen_and_survives_reload(app, client):
    sock = _operator_socket(app)
    try:
        sock.emit("toggle_lyrics_visibility", {"visible": False})
        sock.get_received()
        # Tela sem login lê o estado ao abrir
        assert client.get("/api/lyrics/visible").get_json() == {"visible": False}
        js = client.get("/static/js/pages/bible_display.js").get_data(as_text=True)
        assert "socket.on('lyrics_visibility_changed'" in js and "/api/lyrics/visible" in js
    finally:
        sock.emit("toggle_lyrics_visibility", {"visible": True})
        sock.disconnect()
    assert client.get("/api/lyrics/visible").get_json() == {"visible": True}


def test_new_content_on_screen_shows_hidden_text_again(app, client):
    import display_state

    display_state._live_follow_kind = None
    sock = _operator_socket(app)
    try:
        sock.emit("toggle_lyrics_visibility", {"visible": False})
        sock.get_received()
        # Letra escondida, aí vai um versículo pro telão: aparece
        sock.emit("stage_content", {"kind": "bible", "payload": {"reference": "João 3:16", "text": "..."}, "label": "x"})
        sock.emit("release_staged_content", {"kind": "bible"})
        eventos = [m for m in sock.get_received() if m["name"] == "lyrics_visibility_changed"]
        assert eventos and eventos[0]["args"][0] == {"visible": True}
        assert client.get("/api/lyrics/visible").get_json() == {"visible": True}
    finally:
        sock.emit("toggle_lyrics_visibility", {"visible": True})
        sock.disconnect()
        display_state._live_follow_kind = None


def test_lyrics_and_bible_follow_together_after_both_confirmed(app):
    import display_state
    from extensions import db
    from models import StagedContent

    display_state._live_follow_kind = None
    import telao
    telao.set_public('clock', '/clock')   # letra em lugar nenhum (testes anteriores)
    telao.set_monitor(None)
    sock = _operator_socket(app)

    def preparar(kind, payload):
        sock.emit("stage_content", {"kind": kind, "payload": payload, "label": "x"})
        sock.get_received()
        with app.app_context():
            return db.session.get(StagedContent, kind) is not None   # True = ficou esperando Liberar

    try:
        assert preparar("lyrics", {"songId": 1, "sectionIndex": 0, "title": "V1"})
        sock.emit("release_staged_content", {"kind": "lyrics"})
        # 1ª troca pra Bíblia ainda pede Liberar...
        assert preparar("bible", {"reference": "João 3:16", "text": "..."})
        sock.emit("release_staged_content", {"kind": "bible"})
        sock.get_received()
        # ...depois, voltar pra música e de novo pra Bíblia vai direto
        assert not preparar("lyrics", {"songId": 1, "sectionIndex": 1, "title": "Refrão"})
        assert not preparar("bible", {"reference": "João 3:17", "text": "..."})
        import telao
        assert telao.build_status()["follow_label"] == "Letra e Bíblia"
        # Foto continua pedindo Liberar
        assert preparar("image", {"image_url": "/static/x.jpg"})
        # "parar" volta a pedir pra tudo
        sock.emit("stop_live_follow", {})
        assert preparar("lyrics", {"songId": 1, "sectionIndex": 2, "title": "V2"})
    finally:
        for kind in ("image", "lyrics", "bible"):
            sock.emit("discard_staged_content", {"kind": kind})
        sock.disconnect()
        display_state._live_follow_kind = None


def test_view_song_edit_button_only_for_logged_operator(app, client, logged_in_client):
    """/view/<id> é aberta pros músicos sem login; o Editar só aparece pra
    quem está logado."""
    import songs_store

    with app.app_context():
        songs = songs_store.load_songs()
        song_id = songs_store.get_next_id(songs)
        songs_store.save_songs(songs + [{
            "id": song_id, "title": "Música Editar", "author": "Autor",
            "sections": [{"title": "Verso", "text": ["la la"]}],
        }])
    edit_href = f'href="/edit_song/{song_id}"'
    anon = client.get(f"/view/{song_id}")
    assert anon.status_code == 200 and edit_href not in anon.get_data(as_text=True)
    assert edit_href in logged_in_client.get(f"/view/{song_id}").get_data(as_text=True)


def test_write_routes_require_login(client):
    """Rotas que alteram músicas, a ordem de hoje e o estado das telas
    pediam nada: qualquer um na rede editava ou apagava música."""
    for method, url in [("get", "/edit_song/1"), ("post", "/edit_song/1"), ("post", "/add_song"),
                        ("post", "/delete_song/1"), ("post", "/save"), ("post", "/download_audio"),
                        ("post", "/save_image_path"), ("post", "/save_display_state"),
                        ("post", "/save_lyrics_state")]:
        resp = getattr(client, method)(url)
        assert resp.status_code == 302 and "/login" in resp.headers["Location"], (method, url)


def test_download_audio_only_deletes_inside_audio_folder(app, logged_in_client, tmp_path, monkeypatch):
    import subprocess

    victim = tmp_path / "importante.db"
    victim.write_text("x")
    calls = []

    class Falhou:
        returncode, stderr = 1, "sem rede no teste"

    monkeypatch.setattr(subprocess, "run", lambda cmd, **kw: calls.append(cmd) or Falhou())
    audio_dir = os.path.join(app.static_folder, "audio")
    rel = os.path.relpath(victim, audio_dir)
    resp = logged_in_client.post("/download_audio", json={
        "url": "--exec=touch /tmp/x https://youtu.be/abcdefghijk", "song_id": "../1",
        "old_audio_url": "/static/audio/" + rel,
    })
    assert resp.status_code == 500
    assert victim.exists()  # fora da pasta de áudio: não apaga
    cmd = calls[0]
    assert cmd[-2:] == ["--", "https://www.youtube.com/watch?v=abcdefghijk"]
    assert not any(a.startswith("--exec") for a in cmd)
    assert "/../" not in cmd[cmd.index("-o") + 1]


def test_order_edit_saves_tema_and_horario(app, logged_in_client):
    """/order/<data> manda 'tema' e 'horario'; a rota ignorava os dois."""
    import orders_store

    key = "2031-03-02"
    base = {"tipo": "Culto Dominical", "dirigente": "Ana", "itens": [{"title": "Item"}]}
    assert logged_in_client.post(f"/api/orders/{key}", json={**base, "tema": "Graça", "horario": "18:40"}).status_code == 200
    with app.app_context():
        saved = orders_store.get_order(key)
    assert saved["tema"] == "Graça" and saved["horario"] == "18:40"

    # Editar de novo troca os dois
    logged_in_client.post(f"/api/orders/{key}", json={**base, "tema": "Fé", "horario": "19:00"})
    with app.app_context():
        saved = orders_store.get_order(key)
    assert saved["tema"] == "Fé" and saved["horario"] == "19:00"

    # Quem não manda o campo (ex.: agenda) não apaga o tema
    logged_in_client.post(f"/api/orders/{key}", json=base)
    with app.app_context():
        assert orders_store.get_order(key)["tema"] == "Fé"

    # Apagar o texto do tema apaga do banco
    logged_in_client.post(f"/api/orders/{key}", json={**base, "tema": ""})
    with app.app_context():
        assert "tema" not in orders_store.get_order(key)
    logged_in_client.delete(f"/api/orders/{key}")


def test_new_song_save_and_continue(app, logged_in_client):
    """/new_song tem "Salvar e continuar editando"; /add_song devolve o id
    pra página abrir /edit_song/<id>. A edição não mostra o botão."""
    import songs_store

    assert "btn-salvar-continuar" in logged_in_client.get("/new_song").get_data(as_text=True)
    resp = logged_in_client.post("/add_song", data={
        "title": "Música Continuar", "author": "Autor",
        "sections[]": ["Verso"], "texts[]": ["la la"],
    }).get_json()
    with app.app_context():
        song = songs_store.get_song(resp["id"])
    assert song and song["title"] == "Música Continuar"
    assert "btn-salvar-continuar" not in logged_in_client.get(f"/edit_song/{resp['id']}").get_data(as_text=True)
    logged_in_client.post(f"/delete_song/{resp['id']}")


def test_cifra_club_titulo_e_artista_com_acento():
    """Título e artista vêm do JSON-LD da página (com acento), não do
    endereço ("ministerio-morada" virava "Ministerio Morada")."""
    from bs4 import BeautifulSoup
    from blueprints.cifras import _extrair_artista_cifra, _extrair_titulo_compositor_cifra
    html = """<html><head>
    <script type="application/ld+json">{"@type": ["MusicRecording", "Article"],
      "name": "Ministério Morada - Para Onde Eu Irei?",
      "byArtist": {"@type": "MusicGroup", "name": "Ministério Morada"}}</script>
    <script type="application/ld+json">{"@type": "MusicComposition",
      "name": "Para Onde Eu Irei?", "composer": [{"@type": "Person", "name": "João Sá"}]}</script>
    </head><body><pre data-chord-content="true">C  G</pre></body></html>"""
    soup = BeautifulSoup(html, 'html.parser')
    titulo, compositor, _ = _extrair_titulo_compositor_cifra(soup)
    assert titulo == "Para Onde Eu Irei?"
    assert compositor == "João Sá"
    assert _extrair_artista_cifra(soup) == "Ministério Morada"

    # Sem MusicRecording: usa o 3º item da trilha (BreadcrumbList)
    so_trilha = BeautifulSoup("""<script type="application/ld+json">{"@type": "BreadcrumbList",
      "itemListElement": [{"position": 1, "name": "Início"}, {"position": 2, "name": "Gospel"},
      {"position": 3, "name": "Diante do Trono"}]}</script>""", 'html.parser')
    assert _extrair_artista_cifra(so_trilha) == "Diante do Trono"
    assert _extrair_artista_cifra(BeautifulSoup("<p>nada</p>", 'html.parser')) is None


def test_song_crud_does_not_load_every_song(app, logged_in_client, monkeypatch):
    """Criar, abrir, editar, salvar partitura e excluir mexem só na música
    em questão — load_songs()/save_songs() da lista inteira custavam ~2 s
    por Salvar no Pi, com o servidor parado pra todo mundo nesse tempo."""
    import songs_store

    def proibido(*a, **k):
        raise AssertionError("carregou/gravou todas as músicas")
    import blueprints.songs as rotas
    monkeypatch.setattr(songs_store, "load_songs", proibido)
    monkeypatch.setattr(songs_store, "save_songs", proibido)
    monkeypatch.setattr(rotas, "load_songs", proibido)

    with app.app_context():
        outra = songs_store.next_song_id()
        songs_store.save_song({"id": outra, "title": "Outra", "author": "X",
                               "sections": [{"title": "V", "text": ["intacta"]}]})

    song_id = logged_in_client.post("/add_song", data={
        "title": "Uma Só", "author": "Autor", "sections[]": ["Verso"], "texts[]": ["la la"],
    }).get_json()["id"]
    assert song_id == outra + 1
    assert logged_in_client.get(f"/edit_song/{song_id}").status_code == 200
    resp = logged_in_client.post(f"/edit_song/{song_id}", data={
        "title": "Uma Só (editada)", "author": "Autor",
        "sections[]": ["Refrão"], "texts[]": ["novo texto"],
    })
    assert resp.get_json()["message"] == "Música editada com sucesso!"
    assert logged_in_client.post(f"/api/save_sheet/{song_id}", json={"id": 7, "measures": []}).status_code == 200
    assert logged_in_client.get(f"/sheet_editor/{song_id}").status_code == 200

    with app.app_context():
        song = songs_store.get_song(song_id)
        assert song["title"] == "Uma Só (editada)"
        assert song["sections"] == [{"title": "Refrão", "text": ["novo texto"]}]
        assert [sh["id"] for sh in song["sheet_data"]] == [7]
        assert songs_store.get_song(outra)["sections"] == [{"title": "V", "text": ["intacta"]}]

    logged_in_client.post(f"/delete_song/{song_id}")
    with app.app_context():
        assert songs_store.get_song(song_id) is None
        assert songs_store.get_song(outra) is not None
        songs_store.delete_song(outra)


def test_sheet_marks_round_trip(app):
    """Sinais de repetição (ritornelo, casas, D.C./D.S.) ficam na coluna
    sheet_data.marks e voltam iguais; partitura sem repetição continua sem
    a chave preenchida (None), igual a antes."""
    import songs_store
    marks = {"0": {"start": True}, "1": {"volta": "1", "end": True}, "2": {"volta": "2"}, "3": {"jump": "DC_AL_FINE"}}
    nota = {"pitch": "C4", "duration": "w", "dotted": False, "accidental": "", "rest": False}
    with app.app_context():
        sid = songs_store.next_song_id()
        songs_store.save_song({"id": sid, "title": "Rep", "author": "X", "sections": [],
                               "sheet_data": [{"id": 1, "clef": "treble", "keySig": "C", "timeSig": "4/4",
                                               "measures": [[nota]] * 4, "ties": [], "chords": [], "marks": marks},
                                              {"id": 2, "clef": "treble", "keySig": "C", "timeSig": "4/4",
                                               "measures": [[nota]], "ties": [], "chords": [], "marks": {}}]})
        sheets = songs_store.get_song(sid)["sheet_data"]
        assert sheets[0]["marks"] == marks
        assert sheets[1]["marks"] is None
        songs_store.delete_song(sid)


import pytest  # noqa: E402  (importorskip nos testes de upload abaixo)


def _upload(client, nome, dados, folder=None):
    data = {"images": (dados, nome)}
    if folder:
        data["folder"] = folder
    return client.post("/upload", data=data, content_type="multipart/form-data")


def test_upload_phone_photo_is_rotated_by_exif(logged_in_client, tmp_path, monkeypatch):
    """Foto de celular em pé vem deitada + EXIF orientação 6; o telão
    precisa dela em pé (antes ficava deitada: o EXIF era descartado)."""
    import io
    from PIL import Image

    (tmp_path / "static" / "imagens").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    buf = io.BytesIO()
    exif = Image.Exif()
    exif[0x0112] = 6
    Image.new("RGB", (400, 300), "blue").save(buf, format="JPEG", exif=exif)
    buf.seek(0)

    resp = _upload(logged_in_client, "foto.jpg", buf)
    assert resp.status_code == 200
    with Image.open(tmp_path / "static" / "imagens" / "foto.jpg") as img:
        assert img.size == (300, 400)


def test_upload_heic_and_webp_become_visible_images(logged_in_client, tmp_path, monkeypatch):
    import io
    from PIL import Image
    pytest.importorskip("pillow_heif")

    (tmp_path / "static" / "imagens").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    heic = io.BytesIO()
    Image.new("RGB", (64, 48), "green").save(heic, format="HEIF")
    heic.seek(0)
    webp = io.BytesIO()
    Image.new("RGB", (64, 48), "purple").save(webp, format="WEBP")
    webp.seek(0)

    assert _upload(logged_in_client, "IMG_0001.HEIC", heic).get_json()["saved_files"] == ["IMG_0001.jpg"]
    assert _upload(logged_in_client, "banner.webp", webp).get_json()["saved_files"] == ["banner.jpg"]
    urls = {f["url"] for f in logged_in_client.get("/lista-fotos").get_json()["fotos"]}
    assert urls == {"/static/imagens/IMG_0001.jpg", "/static/imagens/banner.jpg"}


def test_upload_unreadable_file_returns_clear_error(logged_in_client, tmp_path, monkeypatch):
    import io

    (tmp_path / "static" / "imagens").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    resp = _upload(logged_in_client, "planilha.xlsx", io.BytesIO(b"isto nao e imagem"))
    assert resp.status_code == 400
    assert "formato não reconhecido" in resp.get_json()["message"]


def test_upload_pptx_converts_slides_in_background(logged_in_client, tmp_path, monkeypatch):
    import io
    import time
    from blueprints import gallery
    pptx = pytest.importorskip("pptx")
    if not gallery._find_libreoffice():
        pytest.skip("LibreOffice não instalado")

    (tmp_path / "static" / "imagens" / "Avisos").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    # Sem /usr/bin no PATH, como no serviço do Pi
    monkeypatch.setenv("PATH", "/nao/existe")
    pres = pptx.Presentation()
    for titulo in ("Bem-vindos", "Oração"):
        pres.slides.add_slide(pres.slide_layouts[0]).shapes.title.text = titulo
    buf = io.BytesIO()
    pres.save(buf)
    buf.seek(0)

    resp = _upload(logged_in_client, "Avisos do culto.pptx", buf, folder="Avisos")
    assert resp.status_code == 200
    assert resp.get_json()["converting"] == ["Avisos_do_culto.pptx"]

    pasta = tmp_path / "static" / "imagens" / "Avisos"
    for _ in range(240):
        if len(list(pasta.glob("Avisos_do_culto_slide-*.jpg"))) == 2:
            break
        time.sleep(0.5)
    assert sorted(p.name for p in pasta.iterdir()) == [
        "Avisos_do_culto_slide-0001.jpg", "Avisos_do_culto_slide-0002.jpg"]


def _pptx_com_videos(tmp_path, videos):
    """PowerPoint com um slide por item de `videos`: None = sem vídeo,
    ('arquivo', oculto) = vídeo inserido (oculto=True esconde o slide)."""
    import io
    from PIL import Image
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    poster = tmp_path / "poster.png"
    Image.new("RGB", (32, 18), "black").save(poster)
    pres = pptx.Presentation()
    for i, item in enumerate(videos, 1):
        slide = pres.slides.add_slide(pres.slide_layouts[5])
        slide.shapes.title.text = f"Slide {i}"
        if item:
            arquivo, oculto = item
            slide.shapes.add_movie(str(arquivo), Inches(1), Inches(2), Inches(4), Inches(3),
                                   poster_frame_image=str(poster),
                                   mime_type="video/x-ms-wmv" if str(arquivo).endswith(".wmv") else "video/mp4")
            if oculto:
                slide._element.set("show", "0")
    buf = io.BytesIO()
    pres.save(buf)
    buf.seek(0)
    return buf


def test_pptx_videos_mapped_to_exported_page(tmp_path):
    """Vídeo vai pro número da página no PDF (slides ocultos não contam)."""
    from blueprints import gallery

    a, b, c = tmp_path / "a.mp4", tmp_path / "b.mp4", tmp_path / "c.wmv"
    for f in (a, b, c):
        f.write_bytes(b"\x00\x00\x00\x18ftypmp42" + f.name.encode())
    buf = _pptx_com_videos(tmp_path, [None, (a, False), (b, True), (c, False)])
    path = tmp_path / "p.pptx"
    path.write_bytes(buf.read())

    mapa = gallery._pptx_videos_by_slide(str(path))
    assert [(pagina, [m.rsplit(".", 1)[1] for m in midias]) for pagina, midias in mapa] == [(2, ["mp4"]), (3, ["wmv"])]

    destino = tmp_path / "out"
    destino.mkdir()
    import os
    env = dict(os.environ, PATH="/nao/existe")  # sem ffmpeg: o WMV fica de fora, o MP4 vai
    salvos = gallery._extract_pptx_videos(str(path), str(tmp_path), str(destino), "Avisos", env)
    assert salvos == ["Avisos_slide-0002v1.mp4"]
    assert (destino / "Avisos_slide-0002v1.mp4").read_bytes().endswith(b"a.mp4")


def test_monitor_split_pane_matches_content_kind():
    import telao
    telao.set_monitor_split("/img", "/lyrics/display")
    try:
        assert telao.monitor_split_pane("/static/imagens/a.jpg") == "url1"
        assert telao.monitor_split_pane("/static/imagens/v.mp4") == "url1"
        assert telao.monitor_split_pane("/lyrics/display?song_id=1&section=2&host=monitor") == "url2"
        assert telao.monitor_split_pane("/clock") is None       # nenhum lado mostra relógio
        telao.set_monitor_split("/lyrics/display/public", "/clock")
        assert telao.monitor_split_pane("/lyrics/display?song_id=1&section=0") == "url1"
        assert telao.monitor_split_pane("/static/imagens/a.jpg") is None
    finally:
        telao.set_monitor(None)
    assert telao.monitor_split_pane("/static/imagens/a.jpg") is None  # sem divisão


def test_split_monitor_keeps_split_when_content_changes(app):
    """Monitor dividido (foto + cifra): seção nova e foto nova vão pro lado
    certo e a divisão continua — em "Só monitor" a seção ia por
    display_monitor_content e trocava o monitor inteiro pela cifra."""
    import telao
    monitor = ucontrol_main.socketio.test_client(app)  # tela, sem login
    monitor.emit("identify", {"role": "monitor"})
    op = _operator_socket(app)
    try:
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": True})
        op.emit("show_monitor_two_urls", {"url1": "/img", "url2": "/lyrics/display"})
        monitor.get_received()

        op.emit("stage_content", {"kind": "lyrics", "label": "Coro",
                                  "payload": {"songId": 1, "sectionIndex": 2, "title": "Coro", "text": "x", "type": "text"}})
        op.emit("release_staged_content", {"kind": "lyrics"})
        names = [m["name"] for m in monitor.get_received()]
        assert "update_lyrics" in names and "display_monitor_content" not in names
        assert telao._monitor["kind"] == "split"

        op.emit("show_monitor_content", {"url": "/static/imagens/foto.jpg"})
        got = [(m["name"], m["args"][0]) for m in monitor.get_received()]
        assert ("display_image", {"image_url": "/static/imagens/foto.jpg"}) in got
        assert not any(n == "display_monitor_content" for n, _ in got)
        assert telao._monitor["kind"] == "split"

        # "Os dois": foto liberada também mantém a divisão
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.emit("show_image", {"image_url": "/static/imagens/outra.jpg"})
        assert telao._monitor["kind"] == "split"
        # Fonte em tela cheia (relógio) continua saindo da divisão
        op.emit("show_image", {"image_url": "/clock"})
        assert telao._monitor is None
    finally:
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.emit("stop_live_follow", {})
        telao.set_monitor(None)
        op.disconnect()
        monitor.disconnect()


def test_split_monitor_only_mode_skips_release_for_visible_kinds(app):
    """Só monitor + monitor dividido (foto + cifra): alternar entre seção e
    foto não pede Liberar — os dois já estão na tela. Em "Os dois" pede."""
    import telao
    from models import StagedContent
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    op = _operator_socket(app)
    secao = {"kind": "lyrics", "label": "Coro",
             "payload": {"songId": 1, "sectionIndex": 1, "title": "Coro", "text": "x", "type": "text"}}
    foto = {"kind": "image", "label": "Foto", "payload": {"image_url": "/static/imagens/foto.jpg"}}
    try:
        op.emit("stop_live_follow", {})
        import display_state
        display_state._follow_parado = False   # só limpa; o "parar" é testado à parte
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": True})
        op.emit("show_monitor_two_urls", {"url1": "/img", "url2": "/lyrics/display"})
        monitor.get_received()
        for item, evento in ((secao, "update_lyrics"), (foto, "display_image"), (secao, "update_lyrics")):
            op.emit("stage_content", item)
            assert evento in [m["name"] for m in monitor.get_received()]
            with app.app_context():
                assert ucontrol_main.db.session.get(StagedContent, item["kind"]) is None
        # Bíblia não tem lado aqui (e iria pro público): continua pedindo
        op.emit("stage_content", {"kind": "bible", "label": "João 3:16", "payload": {"reference": "João 3:16", "text": "x"}})
        with app.app_context():
            assert ucontrol_main.db.session.get(StagedContent, "bible") is not None

        # "Os dois": a foto já está no lado da galeria — vai direto pra lá
        # (regra "onde a foto já está", _foto_onde_a_foto_esta)
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        monitor.get_received()
        op.emit("stage_content", {**foto, "payload": {"image_url": "/static/imagens/outra.jpg"}})
        with app.app_context():
            assert ucontrol_main.db.session.get(StagedContent, "image") is None
        assert telao._monitor["kind"] == "split"
    finally:
        op.emit("discard_staged_content", {"kind": "image"})
        op.emit("discard_staged_content", {"kind": "bible"})
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.emit("stop_live_follow", {})
        telao.set_monitor(None)
        op.disconnect()
        monitor.disconnect()


def test_public_only_mode_still_moves_lyrics_already_on_monitor(app):
    """Só público: se o monitor já mostra a cifra, a seção nova vai pra ele
    também; se mostra outra coisa (foto), continua congelado."""
    import telao
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    op = _operator_socket(app)
    secao = lambda i: {"kind": "lyrics", "label": "Coro", "payload": {
        "songId": 1, "sectionIndex": i, "title": f"S{i}", "text": "x", "type": "text"}}
    try:
        op.emit("stop_live_follow", {})
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.emit("stage_content", secao(0))
        op.emit("release_staged_content", {"kind": "lyrics"})
        op.emit("set_display_mode", {"exclude_monitor": True, "monitor_only": False})
        monitor.get_received()

        op.emit("stage_content", secao(1))
        op.emit("release_staged_content", {"kind": "lyrics"})
        titulos = [m["args"][0]["title"] for m in monitor.get_received() if m["name"] == "update_lyrics"]
        assert titulos == ["S1"]

        # Monitor com foto (só monitor), depois Só público: fica congelado
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": True})
        op.emit("show_monitor_content", {"url": "/static/imagens/foto.jpg"})
        op.emit("set_display_mode", {"exclude_monitor": True, "monitor_only": False})
        monitor.get_received()
        op.emit("stage_content", secao(2))
        op.emit("release_staged_content", {"kind": "lyrics"})
        assert not [m for m in monitor.get_received() if m["name"] == "update_lyrics"]
        assert telao._monitor["kind"] == "photo"
    finally:
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.emit("stop_live_follow", {})
        telao.set_monitor(None)
        op.disconnect()
        monitor.disconnect()


def test_section_change_goes_where_lyrics_are_shown(app):
    """Seção nova vai direto pra onde a letra já está, não importa o modo:
    público com foto + monitor com a letra = muda só no monitor, a foto
    fica no público. Letra em lugar nenhum = pede Liberar."""
    import display_state
    import telao
    from models import StagedContent
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    publico = ucontrol_main.socketio.test_client(app)
    publico.emit("identify", {"role": "public"})
    op = _operator_socket(app)
    secao = lambda i: {"kind": "lyrics", "label": "Coro", "payload": {
        "songId": 1, "sectionIndex": i, "title": f"S{i}", "text": "x", "type": "text"}}
    def recebeu(cli):
        return [m["args"][0]["title"] for m in cli.get_received() if m["name"] == "update_lyrics"]
    def pendente():
        with app.app_context():
            return ucontrol_main.db.session.get(StagedContent, "lyrics") is not None
    try:
        op.emit("stop_live_follow", {})
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        telao.set_public('clock', '/clock')
        telao.set_monitor(None)
        # Letra em lugar nenhum: pede Liberar
        op.emit("stage_content", secao(0))
        assert pendente()
        op.emit("release_staged_content", {"kind": "lyrics"})   # Os dois: público e monitor
        monitor.get_received(); publico.get_received()

        # Monitor dividido (letra + URL) e público com a letra: muda nos dois
        op.emit("show_monitor_two_urls", {"url1": "/lyrics/display", "url2": "https://exemplo.org"})
        op.emit("stage_content", secao(1))
        assert not pendente() and recebeu(monitor) == ["S1"] and recebeu(publico) == ["S1"]

        # Público passa pra foto (Só público); monitor segue com a letra
        op.emit("set_display_mode", {"exclude_monitor": True, "monitor_only": False})
        op.emit("show_image", {"image_url": "/static/imagens/foto.jpg", "exclude_monitor": True})
        monitor.get_received(); publico.get_received()
        for modo in ({"exclude_monitor": False, "monitor_only": False},
                     {"exclude_monitor": True, "monitor_only": False},
                     {"exclude_monitor": False, "monitor_only": True}):
            op.emit("set_display_mode", modo)
            op.emit("stage_content", secao(2))
            assert not pendente()
            assert recebeu(monitor) == ["S2"] and recebeu(publico) == []   # foto fica no público
            assert telao.build_status()["live"]["kind"] == "photo"

        # "parar" da barrinha: volta a pedir Liberar
        op.emit("stop_live_follow", {})
        op.emit("stage_content", secao(3))
        assert pendente()
    finally:
        op.emit("discard_staged_content", {"kind": "lyrics"})
        op.emit("set_display_mode", {"exclude_monitor": False, "monitor_only": False})
        op.emit("stop_live_follow", {})
        display_state._follow_parado = False
        telao.set_monitor(None)
        for c in (op, monitor, publico):
            c.disconnect()


def test_photo_goes_where_photos_are_shown(app):
    """Foto nova vai direto pra onde já tem foto: público com foto e
    monitor com a letra = só o público muda; monitor com foto e público
    com a letra = só o monitor. "Só público" mantém o monitor congelado.
    Foto em lugar nenhum = pede Liberar."""
    import display_state
    import telao
    from models import StagedContent
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    publico = ucontrol_main.socketio.test_client(app)
    publico.emit("identify", {"role": "public"})
    op = _operator_socket(app)
    foto = lambda n: {"kind": "image", "label": "Foto", "payload": {"image_url": f"/static/imagens/{n}.jpg"}}
    def fotos(cli):
        return [m["args"][0].get("image_url") or m["args"][0].get("url") for m in cli.get_received()
                if m["name"] in ("display_image", "display_monitor_content")]
    def pendente():
        with app.app_context():
            return ucontrol_main.db.session.get(StagedContent, "image") is not None
    def modo(m):
        op.emit("set_display_mode", {"exclude_monitor": m == "public", "monitor_only": m == "monitor"})
    try:
        display_state._follow_parado = False
        display_state._live_follow_kind = None
        modo("all")
        telao.set_public('clock', '/clock')
        telao.set_monitor(None)
        op.emit("stage_content", foto("a"))
        assert pendente()                                  # foto em lugar nenhum
        op.emit("discard_staged_content", {"kind": "image"})

        # Público com foto, monitor com a letra (só monitor)
        telao.set_public('photo', '/static/imagens/a.jpg')
        telao.set_monitor('cifra', '/lyrics/display?song_id=1&section=0&host=monitor')
        display_state._live_follow_kind = None
        monitor.get_received(); publico.get_received()
        for m in ("all", "monitor", "public"):
            modo(m)
            op.emit("stage_content", foto("b"))
            assert not pendente()
            assert fotos(publico) == ["/static/imagens/b.jpg"] and fotos(monitor) == []
            assert telao._monitor["kind"] == "cifra"

        # Monitor com foto, público com a letra
        telao.set_public('lyrics', '/lyrics/display/public')
        telao.set_monitor('photo', '/static/imagens/m.jpg')
        monitor.get_received(); publico.get_received()
        for m in ("all", "monitor"):
            modo(m)
            op.emit("stage_content", foto("c"))
            assert not pendente()
            assert fotos(monitor) == ["/static/imagens/c.jpg"] and fotos(publico) == []
            assert telao.build_status()["live"]["kind"] == "lyrics"
        # "Só público": monitor congelado pras fotos → aqui a foto não tem
        # onde ir, então pede Liberar
        modo("public")
        op.emit("stage_content", foto("d"))
        assert pendente() and fotos(monitor) == []
    finally:
        op.emit("discard_staged_content", {"kind": "image"})
        modo("all")
        display_state._live_follow_kind = None
        telao.set_monitor(None)
        for c in (op, monitor, publico):
            c.disconnect()


def test_bible_goes_only_where_released(app):
    """Versículo: cifra no monitor não vira Bíblia; com Liberar em "Os
    dois" o monitor passa pra Bíblia e os próximos versículos seguem lá.
    "Só público" não mexe no monitor; foto do público não é trocada."""
    import display_state
    import telao
    from models import StagedContent
    monitor = ucontrol_main.socketio.test_client(app)
    monitor.emit("identify", {"role": "monitor"})
    publico = ucontrol_main.socketio.test_client(app)
    publico.emit("identify", {"role": "public"})
    op = _operator_socket(app)
    ver = lambda v: {"kind": "bible", "label": f"João 3:{v}", "payload": {"reference": f"João 3:{v}", "text": "x"}}
    refs = lambda cli: [m["args"][0]["reference"] for m in cli.get_received() if m["name"] == "update_bible"]
    abriu = lambda cli: [m["args"][0]["url"] for m in cli.get_received() if m["name"] == "display_monitor_content"]
    def pendente():
        with app.app_context():
            return ucontrol_main.db.session.get(StagedContent, "bible") is not None
    def modo(m):
        op.emit("set_display_mode", {"exclude_monitor": m == "public", "monitor_only": m == "monitor"})
    try:
        display_state._follow_parado = False
        display_state._live_follow_kind = None
        display_state._last_bible_payload_monitor = None
        modo("all")
        # Público na letra, monitor com a cifra própria
        telao.set_public('lyrics', '/lyrics/display/public')
        telao.set_monitor('cifra', '/lyrics/display?song_id=1&section=0&host=monitor')
        op.emit("stage_content", ver(16))
        assert pendente()                       # 1ª troca música → leitura pede Liberar
        # "Só público": vai só pro público, cifra fica no monitor
        modo("public")
        monitor.get_received(); publico.get_received()
        op.emit("release_staged_content", {"kind": "bible"})
        assert refs(publico) == ["João 3:16"] and refs(monitor) == []
        assert telao._monitor["kind"] == "cifra"
        # Próximo versículo: direto, só no público (cifra continua)
        modo("all")
        op.emit("stage_content", ver(17))
        assert not pendente() and refs(publico) == ["João 3:17"] and refs(monitor) == []

        # Liberar com o monitor incluído ("Os dois"): o monitor abre a Bíblia
        telao.set_public('photo', '/static/imagens/a.jpg')
        display_state._live_follow_kind = None
        op.emit("stage_content", ver(18))
        assert pendente()                       # Bíblia em lugar nenhum
        monitor.get_received(); publico.get_received()
        op.emit("release_staged_content", {"kind": "bible"})
        msgs = monitor.get_received()
        assert any(m["name"] == "display_monitor_content" and m["args"][0]["url"] == "/bible/display" for m in msgs)
        assert refs(publico) == ["João 3:18"]
        assert telao.monitor_mostra_biblia()

        # Público passa pra foto ("Só público"); versículo novo só no monitor
        op.emit("show_image", {"image_url": "/static/imagens/b.jpg", "exclude_monitor": True})
        monitor.get_received(); publico.get_received()
        op.emit("stage_content", ver(19))
        assert not pendente() and refs(monitor) == ["João 3:19"] and refs(publico) == []
        assert telao.build_status()["live"]["kind"] == "photo"
        assert ucontrol_main.app.test_client().get("/api/bible/current?host=monitor").get_json()["reference"] == "João 3:19"

        # Volta pra música no público ("Os dois"): o monitor sai da Bíblia
        # e volta pra cifra
        telao.set_public('lyrics', '/lyrics/display/public')
        display_state._live_follow_kind = 'lyrics'
        monitor.get_received()
        op.emit("stage_content", {"kind": "lyrics", "label": "Coro", "payload": {
            "songId": 1, "sectionIndex": 3, "title": "Coro", "text": "x", "type": "text"}})
        urls = [m["args"][0]["url"] for m in monitor.get_received() if m["name"] == "display_monitor_content"]
        assert urls == ["/lyrics/display?song_id=1&section=3&host=monitor"]
        assert telao._monitor["kind"] == "cifra"
    finally:
        op.emit("discard_staged_content", {"kind": "bible"})
        modo("all")
        display_state._live_follow_kind = None
        telao.set_monitor(None)
        for c in (op, monitor, publico):
            c.disconnect()


def test_cifra_import_only_fetches_cifra_club(app, client, logged_in_client, monkeypatch):
    """/cifra e /cifra_com_video buscavam QUALQUER endereço, sem login:
    dava pra usar o Pi pra acessar a rede interna. Agora exigem login,
    só aceitam cifraclub.com.br e conferem o redirecionamento."""
    import requests
    pedidos = []

    class Resp:
        def __init__(self, status=200, location=None, text="<html></html>"):
            self.status_code, self.text = status, text
            self.headers = {"Location": location} if location else {}
            self.is_redirect = location is not None

    respostas = {}
    def fake_get(url, **kw):
        pedidos.append(url)
        assert kw.get("timeout") and kw.get("allow_redirects") is False
        return respostas.get(url, Resp())
    monkeypatch.setattr(requests, "get", fake_get)

    assert client.get("/cifra?url=https://www.cifraclub.com.br/x/y/").status_code in (302, 401)  # exige login
    for rota in ("/cifra", "/cifra_com_video"):
        for ruim in ("http://192.168.0.1/", "http://127.0.0.1:8585/", "https://cifraclub.com.br.evil.com/",
                     "file:///etc/passwd", "https://user@www.cifraclub.com.br/", "https://www.cifraclub.com.br:8080/"):
            assert logged_in_client.get(rota, query_string={"url": ruim}).status_code == 400, (rota, ruim)
    assert pedidos == []    # nada disso foi buscado

    assert logged_in_client.get("/cifra_com_video", query_string={"url": "https://www.cifraclub.com.br/a/b/"}).status_code == 200
    assert logged_in_client.get("/cifra", query_string={"url": "https://m.cifraclub.com.br/a/b/"}).status_code == 200

    # Redirecionamento pra fora do Cifra Club é barrado (sem buscar o destino)
    pedidos.clear()
    respostas["https://www.cifraclub.com.br/r/"] = Resp(301, "http://10.0.0.1/admin")
    assert logged_in_client.get("/cifra", query_string={"url": "https://www.cifraclub.com.br/r/"}).status_code == 400
    assert pedidos == ["https://www.cifraclub.com.br/r/"]


def test_limpar_route_removed(client):
    assert client.get("/limpar").status_code == 404


def test_image_zoom_survives_screen_reload(app, client):
    """Zoom da foto fica guardado no servidor: a tela que recarrega pede em
    /api/image-zoom e reaplica. Outra foto na tela apaga o zoom."""
    import display_state
    op = _operator_socket(app)
    try:
        display_state._image_zoom = None
        op.emit("show_image", {"image_url": "/static/imagens/a.jpg"})
        op.emit("show_image_zoom", {"image_url": "/static/imagens/a.jpg", "zoom": 2, "pan": {"x": 10, "y": -5}})
        z = client.get("/api/image-zoom").get_json()   # tela sem login
        assert z["image_url"] == "/static/imagens/a.jpg" and z["zoom"] == 2 and z["pan"] == {"x": 10, "y": -5}
        # O zoom vem JUNTO com a foto em tudo que diz à tela o que mostrar —
        # a tela já desenha com zoom (sem a foto normal antes)
        assert client.get("/last_image").get_json()["zoom"]["zoom"] == 2
        assert client.get("/api/current-display").get_json()["zoom"]["zoom"] == 2
        # Depois de reiniciar o servidor o "modo atual" em memória some e a
        # foto vem do banco (DisplayState) — o zoom tem que vir junto também
        import time as _t
        modo, cache = display_state._current_display_mode, dict(display_state._display_state_cache)
        display_state._current_display_mode = None
        display_state._display_state_cache['data'] = None
        _t.sleep(0.3)   # gravação do DisplayState roda numa thread
        try:
            for rota in ("/api/current-display", "/get_display_state", "/last_image"):
                r = client.get(rota).get_json()
                assert r.get("zoom", {}).get("zoom") == 2, (rota, r)
        finally:
            display_state._current_display_mode = modo
            display_state._display_state_cache.update(cache)
        op.get_received()
        op.emit("show_image", {"image_url": "/static/imagens/a.jpg?t=123"})   # mesma foto de novo
        enviado = [m["args"][0] for m in op.get_received() if m["name"] == "display_image"]
        assert enviado and enviado[-1]["zoom"]["zoom"] == 2
        assert client.get("/api/image-zoom").get_json()["zoom"] == 2
        op.emit("show_image", {"image_url": "/static/imagens/b.jpg"})         # outra foto
        assert client.get("/api/image-zoom").get_json() == {}
        assert "zoom" not in client.get("/last_image").get_json()
    finally:
        display_state._image_zoom = None
        op.disconnect()


def test_data_dir_keeps_state_out_of_code_folder(tmp_path):
    """UCONTROL_DATA_DIR (instalação Windows, ver paths.py): banco, mídia e
    QR codes vão pra pasta de dados, servidos nas mesmas URLs /static/...
    Roda num processo à parte porque paths.py lê a variável na importação."""
    import subprocess
    import sys

    script = (
        "import os, main\n"
        "d = os.environ['UCONTROL_DATA_DIR']\n"
        "open(os.path.join(d, 'imagens', 'x.png'), 'wb').write(b'png')\n"
        "c = main.app.test_client()\n"
        "assert c.get('/static/imagens/x.png').data == b'png'\n"
        "assert c.get('/static/imagens/../instance/secret_key').status_code == 404\n"
        "assert c.get('/static/qrcodes/control_qr_code.png').status_code == 200\n"
        "assert c.get('/static/js/telao_bar.js').status_code == 200\n"
        "assert c.get('/imagens').get_json() == ['/static/imagens/x.png']\n"
        "assert main.app.instance_path == os.path.join(d, 'instance')\n"
        "from blueprints.bible import load_bible_metadata\n"
        "assert load_bible_metadata()['books']\n"
    )
    env = {k: v for k, v in os.environ.items() if k != "UCONTROL_TESTING"}
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    env.update(UCONTROL_DATA_DIR=str(tmp_path), UCONTROL_NO_BROWSER="1", PYTHONPATH=root)
    # cwd fora do código: o atalho do Windows não abre na pasta do programa
    result = subprocess.run([sys.executable, "-c", script], cwd=tmp_path, env=env,
                            capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stderr[-2000:]
    assert (tmp_path / "instance" / "secret_key").exists()


def test_first_run_creates_first_user_only_once(app):
    """Instalação nova não tem usuário e /register não existe mais: a tela
    de login oferece criar o primeiro, e só enquanto não houver nenhum."""
    with app.app_context():
        User.query.delete()
        ucontrol_main.db.session.commit()
    try:
        c = app.test_client()
        assert "Repita a senha" in c.get("/login").get_data(as_text=True)

        c.post("/primeiro-acesso", data={"username": "op", "password": "a", "password_confirm": "b"})
        with app.app_context():
            assert User.query.count() == 0

        r = c.post("/primeiro-acesso", data={"username": "op", "password": "a", "password_confirm": "a"})
        assert r.status_code == 302 and r.headers["Location"].endswith("/dashboard")
        assert c.get("/dashboard").status_code == 200

        outro = app.test_client()
        assert "Repita a senha" not in outro.get("/login").get_data(as_text=True)
        outro.post("/primeiro-acesso", data={"username": "intruso", "password": "x", "password_confirm": "x"})
        with app.app_context():
            assert [u.username for u in User.query.all()] == ["op"]
    finally:
        with app.app_context():
            User.query.filter_by(username="op").delete()
            ucontrol_main.db.session.commit()


def test_copyrighted_bibles_only_from_private_folder(tmp_path, monkeypatch):
    """Repositório/instalador só traz as Bíblias de licença aberta; NVI, ACF
    etc. entram pela pasta privada (fora do git). Sem a ACF, o padrão é a
    Bíblia Livre; nome de tradução vindo da URL não sai da pasta."""
    import paths
    from blueprints import bible

    monkeypatch.setattr(paths, "BIBLES_PRIVATE_DIR", str(tmp_path))
    monkeypatch.setattr(bible, "_bible_metadata", None)
    monkeypatch.setattr(bible, "_bible_cache", {})
    meta = bible.load_bible_metadata()
    assert "blivre" in meta["translations"] and "nvi" not in meta["translations"]
    assert not any(n.startswith("nvi") for n in os.listdir(paths.BIBLES_DIR))
    assert bible.default_translation() == "blivre"
    assert bible._bible_path("../instance/bibles/x") is None

    (tmp_path / "nvi.json").write_text('[{"abbrev": "gn", "chapters": [["No princípio"]]}]', encoding="utf-8")
    monkeypatch.setattr(bible, "_bible_metadata", None)
    assert "nvi" in bible.load_bible_metadata()["translations"]
    assert bible.load_bible("nvi")[0]["chapters"][0][0] == "No princípio"
    monkeypatch.setattr(bible, "_bible_metadata", None)


def test_help_has_every_menu_section_and_no_default_password(client):
    """Ajuda abre sem login (e tem QR code): já ensinou uma senha padrão
    admin/1234 e tinha contato falso. Todo tópico do menu tem conteúdo."""
    import re

    html = client.get("/help").get_data(as_text=True)
    js = open(os.path.join(ucontrol_main.app.static_folder, "js", "pages", "help.js"), encoding="utf-8").read()
    secoes = re.findall(r'data-section="(\w+)"', html)
    assert secoes and all(re.search(rf"^    {s}: `", js, re.M) for s in secoes)
    assert "1234" not in js and "exemplo.com" not in js and "imges_help" not in js


def test_cannot_remove_self_or_last_user(app, logged_in_client):
    """Sem nenhum usuário, /login volta a oferecer "criar o primeiro" pra
    quem chegar antes — então ninguém remove a si mesmo nem o último."""
    from werkzeug.security import generate_password_hash

    with app.app_context():
        eu = User.query.filter_by(username="tester").first()
        outro = User(username="outro", password=generate_password_hash("x", method="pbkdf2:sha256"))
        ucontrol_main.db.session.add(outro)
        ucontrol_main.db.session.commit()
        eu_id, outro_id = eu.id, outro.id

    logged_in_client.post(f"/remover_usuario/{eu_id}")
    with app.app_context():
        assert ucontrol_main.db.session.get(User, eu_id) is not None

    logged_in_client.post(f"/remover_usuario/{outro_id}")
    with app.app_context():
        assert ucontrol_main.db.session.get(User, outro_id) is None
        assert User.query.count() >= 1


def test_update_check_compares_versions_and_reads_release(monkeypatch):
    import updates

    assert updates.parse_version("v1.10.0") > updates.parse_version("1.9.9")
    assert updates.parse_version("lixo") < updates.parse_version("0.0.1")

    class Resp:
        status_code = 200

        def raise_for_status(self):
            pass

        def json(self):
            return {"tag_name": "v9.9.9", "body": "Novidades", "html_url": "https://github.com/x",
                    "assets": [{"name": "U-Control-Setup-9.9.9.exe", "size": 3, "digest": "sha256:ab",
                                "browser_download_url": updates.DOWNLOAD_PREFIX + "v9.9.9/U-Control-Setup-9.9.9.exe"}]}

    monkeypatch.setattr(updates.requests, "get", lambda *a, **kw: Resp())
    rel = updates._fetch_latest()
    assert rel["version"] == "9.9.9" and rel["installer"]["digest"] == "sha256:ab"

    Resp.status_code = 404  # repositório ainda sem Release: consultou, não há versão
    assert updates._fetch_latest() is None


def test_update_status_and_install_only_from_the_server_pc(logged_in_client, client, monkeypatch):
    """Selo/página de versão nova; "Baixar e instalar" (fecha o programa)
    só no Windows e só pra quem está no próprio PC — nunca pelo celular
    nem pelo túnel (que também chega de 127.0.0.1)."""
    import updates

    rel = {"version": "9.9.9", "notes": "Novidades", "url": "https://github.com/x",
           "installer": {"name": "s.exe", "url": "https://exemplo.com/s.exe", "size": 1, "digest": None}}
    monkeypatch.setitem(updates._cache, "release", rel)
    monkeypatch.setitem(updates._cache, "quando", 10**12)  # não busca no GitHub

    assert client.get("/api/update-status").status_code == 302  # exige login
    s = logged_in_client.get("/api/update-status").get_json()
    assert s["available"] and s["current"] == updates.current_version() and s["kind"] == "servidor"
    assert not s["can_install"]
    assert 'id="site-nav-update"' in logged_in_client.get("/dashboard").get_data(as_text=True)
    assert logged_in_client.get("/atualizacao").status_code == 200

    monkeypatch.setattr(updates, "install_kind", lambda: "windows")
    assert logged_in_client.get("/api/update-status").get_json()["can_install"]
    celular = {"REMOTE_ADDR": "192.168.0.20"}
    assert not logged_in_client.get("/api/update-status", environ_base=celular).get_json()["can_install"]
    assert logged_in_client.post("/api/update/install", environ_base=celular).status_code == 403
    assert logged_in_client.post("/api/update/install", headers={"CF-Connecting-IP": "1.2.3.4"}).status_code == 403
    # instalador fora das Releases do repositório: recusa, não baixa
    r = logged_in_client.post("/api/update/install")
    assert r.status_code == 400 and updates._instalacao["estado"] is None

    monkeypatch.setitem(rel, "version", "0.0.1")
    assert not logged_in_client.get("/api/update-status").get_json()["available"]


def test_libreoffice_found_in_windows_program_files(monkeypatch, tmp_path):
    """O instalador do LibreOffice no Windows não põe no PATH: antes o
    U-Control nunca achava, mesmo instalado."""
    from blueprints import gallery

    soffice = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    soffice.parent.mkdir(parents=True)
    soffice.write_text("")
    monkeypatch.setattr(gallery.shutil, "which", lambda name: None)
    monkeypatch.setenv("ProgramFiles", str(tmp_path))
    real_exists = os.path.exists
    monkeypatch.setattr(gallery.os.path, "exists", lambda p: p == str(soffice) or (real_exists(p) and not p.startswith("/usr") and not p.startswith("/Applications")))
    assert gallery._find_libreoffice() == str(soffice)


def test_ticker_each_screen_keeps_its_own_aviso(app):
    """Mandar pra uma tela não mexe nas outras: aviso A no público, B no
    palco, nada na transmissão; tirar do palco deixa o do público."""
    op = _operator_socket(app)
    try:
        op.emit("stop_ticker", {"target": "all"})
        op.emit("set_ticker", {"text": "A", "targets": ["public"]})
        op.emit("set_ticker", {"text": "B", "targets": ["monitor"]})
        api = app.test_client().get("/api/ticker").get_json()
        assert api["by_target"] == {"public": "A", "monitor": "B"}

        op.emit("stop_ticker", {"targets": ["monitor"]})
        assert app.test_client().get("/api/ticker").get_json()["by_target"] == {"public": "A"}

        op.emit("set_ticker", {"text": "C", "targets": ["public", "obs"]})
        assert app.test_client().get("/api/ticker").get_json()["by_target"] == {"public": "C", "obs": "C"}

        op.emit("set_ticker", {"text": "Todos", "target": "all"})  # formato antigo
        assert set(app.test_client().get("/api/ticker").get_json()["by_target"]) == {"public", "monitor", "obs"}
    finally:
        op.emit("stop_ticker", {"target": "all"})
        op.disconnect()


# --- PowerPoint animado (blueprints/apresentacao.py) ---------------------

def _pptx_simples(tmp_path, nome, oculto=None, transicao=False):
    import io
    pptx = pytest.importorskip("pptx")
    from lxml import etree
    pres = pptx.Presentation()
    for i in range(3):
        slide = pres.slides.add_slide(pres.slide_layouts[5])
        slide.shapes.title.text = f"Slide {i + 1}"
        if i == oculto:
            slide._element.set("show", "0")
        if transicao and i == 1:
            slide._element.append(etree.fromstring(
                '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:fade/></p:transition>'))
    path = tmp_path / nome
    buf = io.BytesIO()
    pres.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def test_pptx_animado_detecta_transicao_e_tira_ocultos(tmp_path):
    import zipfile
    from blueprints import gallery

    assert not gallery._pptx_animado(str(_pptx_simples(tmp_path, "parado.pptx")))
    animado = _pptx_simples(tmp_path, "animado.pptx", oculto=0, transicao=True)
    assert gallery._pptx_animado(str(animado))

    # SVG do LibreOffice inclui slide oculto e o PDF não: a cópia sem os
    # ocultos mantém o slide N do SVG = imagem N da galeria
    copia = gallery._pptx_sem_ocultos(str(animado), str(tmp_path))
    with zipfile.ZipFile(copia) as zf:
        assert zf.read("ppt/presentation.xml").decode().count("<p:sldId ") == 2
    sem_oculto = _pptx_simples(tmp_path, "todos.pptx", transicao=True)
    assert gallery._pptx_sem_ocultos(str(sem_oculto), str(tmp_path)) == str(sem_oculto)


def test_lista_fotos_marca_slides_animados_e_exclusao_leva_o_svg(logged_in_client, tmp_path, monkeypatch):
    pasta = tmp_path / "static" / "imagens" / "Culto"
    pasta.mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    for nome in ("Avisos_slide-0001.jpg", "Avisos_slide-0002.jpg", "foto.jpg", "Outro_slide-0001.jpg"):
        (pasta / nome).write_bytes(b"jpg")
    (pasta / "Avisos_slide.svg").write_text("<svg/>")

    fotos = {f["url"].rsplit("/", 1)[1]: f for f in logged_in_client.get("/lista-fotos?folder=Culto").get_json()["fotos"]}
    assert fotos["Avisos_slide-0002.jpg"]["animado"] == {"svg": "/static/imagens/Culto/Avisos_slide.svg", "slide": 1}
    assert "animado" not in fotos["foto.jpg"]
    assert "animado" not in fotos["Outro_slide-0001.jpg"]  # sem SVG ao lado
    assert "Avisos_slide.svg" not in fotos  # o SVG não aparece como foto

    logged_in_client.post("/delete_images", json={"folder": "Culto", "image_names": ["Avisos_slide-0001.jpg"]})
    assert (pasta / "Avisos_slide.svg").exists()  # ainda tem um slide
    logged_in_client.post("/delete_images", json={"folder": "Culto", "image_names": ["Avisos_slide-0002.jpg"]})
    assert not (pasta / "Avisos_slide.svg").exists()


def test_apresentacao_sincroniza_telas(app, client):
    from urllib.parse import quote
    import telao
    from blueprints import apresentacao

    svg = "/static/imagens/Culto/Avisos_slide.svg"
    url = lambda n: f"/apresentacao?svg={quote(svg)}&slide={n}"
    assert client.get(url(0)).status_code == 200  # tela: sem login

    tela = ucontrol_main.socketio.test_client(app)  # sem login
    operador = _operator_socket(app)
    try:
        operador.emit("show_image", {"image_url": url(1)})
        assert client.get("/api/apresentacao").get_json()["slide"] == 1
        assert telao.describe({"kind": "page", "url": url(1)})["label"] == "Apresentação · Avisos"
        tela.get_received()

        # a tela conta onde parou (sem login); o servidor guarda e repassa
        tela.emit("apresentacao_pos", {"svg": svg, "slide": 1, "step": 1, "slides": 3, "steps": 2})
        e = client.get("/api/apresentacao").get_json()
        assert (e["slide"], e["step"], e["slides"]) == (1, 1, 3)
        tela.emit("apresentacao_pos", {"svg": "/static/imagens/x_slide.svg", "slide": 9, "step": 0, "slides": 9, "steps": 0})
        assert client.get("/api/apresentacao").get_json()["slide"] == 1  # outra apresentação: ignora

        # Avançar vai pra todas com a posição de partida
        tela.get_received()
        operador.emit("apresentacao_cmd", {"cmd": "next"})
        cmds = [m["args"][0] for m in tela.get_received() if m["name"] == "apresentacao_cmd"]
        assert cmds == [{"cmd": "next", "svg": svg, "slide": 1, "step": 1}]

        # tela sem login não comanda
        tela.emit("apresentacao_cmd", {"cmd": "next"})
        assert not [m for m in tela.get_received() if m["name"] == "apresentacao_cmd"]

        # outro slide da mesma apresentação: pula sem recarregar as telas
        operador.emit("show_image", {"image_url": url(2)})
        recebidos = tela.get_received()
        assert not [m for m in recebidos if m["name"] == "display_image"]
        assert [m["args"][0] for m in recebidos if m["name"] == "apresentacao_cmd"] == [
            {"cmd": "goto", "svg": svg, "slide": 2, "step": 0}]
    finally:
        tela.disconnect()
        operador.disconnect()
        apresentacao._estado.update(svg=None, slide=0, step=0, slides=None, steps=None)


def test_apresentacao_parse_url_recusa_caminho_de_fora():
    from blueprints.apresentacao import parse_url
    assert parse_url("/apresentacao?svg=/static/imagens/a_slide.svg&slide=3") == ("/static/imagens/a_slide.svg", 3)
    assert parse_url("/apresentacao?svg=/etc/passwd&slide=1")[0] is None
    assert parse_url("/apresentacao?svg=/static/imagens/a.jpg")[0] is None
    assert parse_url("/clock") == (None, 0)


# --- Backup (blueprints/backup.py) ----------------------------------------

def test_backup_exige_login(client):
    for url in ("/backup", "/api/backup/baixar", "/api/backup/resumo"):
        assert client.get(url).status_code in (302, 401)
    assert client.post("/api/backup/restaurar").status_code in (302, 401)


def test_backup_baixar_traz_manifest_banco_e_fotos(logged_in_client, tmp_path, monkeypatch):
    import io
    import json
    import zipfile
    import paths

    (tmp_path / "static" / "imagens" / "Culto").mkdir(parents=True)
    (tmp_path / "static" / "imagens" / "Culto" / "foto.jpg").write_bytes(b"jpg")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(paths, "AUDIO_DIR", str(tmp_path / "audio"))
    monkeypatch.setattr(paths, "BIBLES_PRIVATE_DIR", str(tmp_path / "bibles"))

    resp = logged_in_client.get("/api/backup/baixar")
    assert resp.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(resp.data))
    assert {"manifest.json", "banco/users.db", "imagens/Culto/foto.jpg"} <= set(zf.namelist())
    assert json.loads(zf.read("manifest.json"))["app"] == "U-Control"


def test_backup_restaurar_recusa_arquivo_estranho(logged_in_client):
    import io
    import json
    import zipfile

    def enviar(dados):
        return logged_in_client.post("/api/backup/restaurar", data={"arquivo": (io.BytesIO(dados), "b.zip")},
                                     content_type="multipart/form-data")

    assert "não é um backup" in enviar(b"qualquer coisa").get_json()["error"]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("manifest.json", json.dumps({"app": "U-Control", "formato": 1}))
        zf.writestr("banco/users.db", b"")
        zf.writestr("imagens/../../fora.txt", "x")
    resp = enviar(buf.getvalue())
    assert resp.status_code == 400 and "estranho" in resp.get_json()["error"]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("manifest.json", json.dumps({"app": "U-Control", "formato": 99}))
        zf.writestr("banco/users.db", b"")
    assert "mais nova" in enviar(buf.getvalue()).get_json()["error"]


# --- Exportar e importar (blueprints/pacote.py) ---------------------------

def test_pacote_mesma_musica_ignora_maiuscula_acento_e_espaco():
    from blueprints.pacote import _chave
    assert _chave("Grande é o Senhor", "Adhemar ") == _chave("grande e o  senhor", "ADHEMAR")
    assert _chave("Santo", "A") != _chave("Santo", "B")


def test_pacote_reescreve_links_do_culto_pros_ids_daqui():
    from blueprints.pacote import _reescrever_ordem
    ordem = {"tipo": "Culto", "itens": [
        {"title": "a", "url": "/lyrics/3"}, {"title": "b", "url": "/lyrics/30?x=1"},
        {"title": "c", "url": "/order/2026-10-25"}, {"anyText": "Avisos"}]}
    nova = _reescrever_ordem(ordem, {3: 101, 30: 7})
    assert [i.get("url") for i in nova["itens"]] == ["/lyrics/101", "/lyrics/7?x=1", "/order/2026-10-25", None]
    assert ordem["itens"][0]["url"] == "/lyrics/3"  # não mexe no original


def test_pacote_recusa_backup_e_caminho_malicioso(tmp_path):
    import json
    import zipfile
    from blueprints.pacote import PacoteInvalido, ler_pacote

    def zipar(nome, manifest, extra=()):
        caminho = tmp_path / nome
        with zipfile.ZipFile(caminho, "w") as zf:
            zf.writestr("manifest.json", json.dumps(manifest))
            zf.writestr("pacote.json", json.dumps({"musicas": [], "cultos": {}, "pastas": ["Fotos"]}))
            for n in extra:
                zf.writestr(n, "x")
        return str(caminho)

    with pytest.raises(PacoteInvalido, match="backup completo"):
        ler_pacote(zipar("b.zip", {"app": "U-Control", "formato": 1}))
    with pytest.raises(PacoteInvalido, match="estranho"):
        ler_pacote(zipar("m.zip", {"app": "U-Control", "tipo": "pacote", "formato": 1}, ["imagens/Fotos/../../x.jpg"]))
    with pytest.raises(PacoteInvalido, match="estranho"):
        ler_pacote(zipar("o.zip", {"app": "U-Control", "tipo": "pacote", "formato": 1}, ["imagens/Outra/x.jpg"]))
    _, dados, arquivos = ler_pacote(zipar("ok.zip", {"app": "U-Control", "tipo": "pacote", "formato": 1}, ["imagens/Fotos/a.jpg"]))
    assert arquivos == {"Fotos": [("a.jpg", 1)]}


def test_pacote_paginas_exigem_login(client):
    for url in ("/pacote", "/api/pacote/itens"):
        assert client.get(url).status_code in (302, 401)
    for url in ("/api/pacote/exportar", "/api/pacote/analisar", "/api/pacote/importar"):
        assert client.post(url).status_code in (302, 401)


# --- Desligar o U-Control (blueprints/desligar.py, só Windows) -----------

def test_desligar_so_aparece_no_windows_e_no_proprio_pc(logged_in_client, monkeypatch):
    import desktop_ponte

    # Pi/Docker: ninguém registrou -> sem botão e sem rota
    monkeypatch.setattr(desktop_ponte, "_desligar", None)
    assert "Desligar o U-Control" not in logged_in_client.get("/control").get_data(as_text=True)
    assert logged_in_client.get("/desligar").status_code == 404
    assert logged_in_client.post("/api/desligar").status_code == 403

    chamado = []
    monkeypatch.setattr(desktop_ponte, "_desligar", lambda: chamado.append(1))
    monkeypatch.setattr(desktop_ponte.threading, "Timer", lambda _s, f: type("T", (), {"start": staticmethod(f)}))
    # pelo túnel / celular: não
    tunel = {"X-Forwarded-For": "200.1.2.3"}
    assert "Desligar o U-Control" not in logged_in_client.get("/control", headers=tunel).get_data(as_text=True)
    assert logged_in_client.post("/api/desligar", headers=tunel).status_code == 403
    assert logged_in_client.post("/api/desligar", environ_base={"REMOTE_ADDR": "192.168.0.20"}).status_code == 403
    assert not chamado
    # no próprio computador: sim
    assert "Desligar o U-Control" in logged_in_client.get("/control").get_data(as_text=True)
    assert logged_in_client.get("/desligar").status_code == 200
    assert logged_in_client.post("/api/desligar").get_json() == {"ok": True}
    assert chamado == [1]


def test_desligar_exige_login(client, monkeypatch):
    import desktop_ponte
    monkeypatch.setattr(desktop_ponte, "_desligar", lambda: None)
    assert client.post("/api/desligar").status_code in (302, 401)


# --- Fontes e imagens do PowerPoint (pptx_fontes.py, pptx_imagens.py) ----

def test_fonte_embutida_sai_do_eot():
    import struct
    import pptx_fontes

    ttf = b"\x00\x01\x00\x00" + b"fonte de verdade"
    def eot(flags, dados):
        cabecalho = bytearray(82)
        struct.pack_into("<IIII", cabecalho, 0, len(cabecalho) + len(dados), len(dados), 0x20002, flags)
        struct.pack_into("<H", cabecalho, 34, 0x504C)
        return bytes(cabecalho) + dados

    assert pptx_fontes._ttf_do_eot(eot(0, ttf)) == ttf
    assert pptx_fontes._ttf_do_eot(eot(0x10000000, bytes(b ^ 0x50 for b in ttf))) == ttf  # XOR
    assert pptx_fontes._ttf_do_eot(eot(0x4, ttf)) is None  # comprimida (MTX): não dá
    assert pptx_fontes._ttf_do_eot(ttf) == ttf  # TTF puro
    assert pptx_fontes._ttf_do_eot(b"qualquer coisa") is None


def test_gif_girado_volta_a_girar_no_svg():
    import pptx_imagens

    # GIF do Canva girado -111,9°: o LibreOffice grava a caixa do giro
    # (x=-27189 y=-11318 46238x59704) com o GIF reto e esticado dentro
    pic = ('<p:pic><a:blip r:embed="rId2"/><a:xfrm flipH="false" flipV="false" rot="-6714334">'
           '<a:off x="-10978702" y="1527947"/><a:ext cx="19028773" cy="10288223"/></a:xfrm></p:pic>')

    class Zip:
        def read(self, _):
            return pic.encode()

    girados = pptx_imagens._gifs_girados(Zip(), "s", {"rId2": "ppt/media/a.gif"}, {"ppt/media/a.gif"})
    attrs = pptx_imagens._girar(' x="-27189" y="-11318" width="46238" height="59704" ', girados)
    assert 'transform="rotate(-111.906' in attrs
    assert 'width="52858" height="28578"' in attrs  # tamanho reto do PowerPoint
    sem = ' x="0" y="0" width="100" height="100" '
    assert pptx_imagens._girar(sem, girados) == sem  # outra imagem: não mexe


def test_conversoes_mostram_a_etapa_por_pasta(client):
    import conversoes
    conversoes.comecar("Culto", "Avisos", animada=True)
    conversoes.comecar("Outra", "Fotos", animada=False)
    assert [c["etapa"] for c in client.get("/api/conversoes?folder=Culto").get_json()] == ["slides"]
    conversoes.etapa("Culto", "Avisos", "animada", slides=10)
    c = client.get("/api/conversoes?folder=Culto").get_json()[0]
    assert (c["nome"], c["etapa"], c["slides"]) == ("Avisos", "animada", 10)
    conversoes.etapa("Culto", "Avisos", "pronto", animada=True)
    conversoes.etapa(None, "Avisos", "erro")  # conversão sem registro: ignora
    assert client.get("/api/conversoes?folder=Culto").get_json()[0]["etapa"] == "pronto"
    conversoes._estado.clear()


# --- Importar músicas (importar_musicas.py, blueprints/importar.py) ------

def test_importar_le_os_formatos_e_poe_acorde_em_cima_da_letra():
    import io
    import json
    import zipfile
    import importar_musicas as im

    assert im._linha_com_acordes("[G]Santo, [D/F#]santo") == "G      D/F#\nSanto, santo"
    assert im._linha_com_acordes("sem acorde") == "sem acorde"

    cho = im.ler("x.cho", b"{title: Santo}\n{artist: Fulano}\n[G]Quem [D]vem\n\n{soc}\n[C]Coro\n{eoc}\n")[0]
    assert (cho["title"], cho["author"]) == ("Santo", "Fulano")
    assert [s["title"] for s in cho["sections"]] == ["Verse 1", "Chorus"]
    assert cho["sections"][0]["text"] == ["G    D\nQuem vem"]

    ol = ('<song xmlns="http://openlyrics.info/namespace/2009/song"><properties><titles><title>Graça</title></titles>'
          '<authors><author>Newton</author></authors></properties><lyrics><verse name="v1"><lines>'
          '<chord name="D"/>Graça <chord name="G"/>sublime<br/>linha dois</lines></verse><verse name="c">'
          '<lines>Coro</lines></verse></lyrics></song>').encode()
    m = im.ler("g.xml", ol)[0]
    assert [s["title"] for s in m["sections"]] == ["Verse 1", "Chorus"]
    assert m["sections"][0]["text"] == ["D     G\nGraça sublime\nlinha dois"]

    holy = json.dumps([{"title": "Hino", "artist": "Harpa", "slides": [{"text": "a\nb", "slide_description": "Coro"}]},
                       {"title": "Antigo", "lyrics": {"paragraphs": [{"text": "x"}]}}]).encode()
    h = im.ler("h.json", holy)
    assert [(x["title"], x["sections"][0]["title"]) for x in h] == [("Hino", "Coro"), ("Antigo", "Verse 1")]

    t = im.ler("Hino da Casa.txt", "Primeiro\n\n[Refrão]\nGlória\n".encode("cp1252"))[0]
    assert t["title"] == "Hino da Casa" and [s["title"] for s in t["sections"]] == ["Verse 1", "Refrão"]

    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("a.cho", "{title: A}\nla")
        zf.writestr("b.txt", "lb")
        zf.writestr("c.muf", "?")
    assert sorted(m["title"] for m in im.ler("lote.zip", z.getvalue())) == ["A", "b"]
    with pytest.raises(im.ArquivoInvalido, match="muf"):
        im.ler("x.muf", b"?")
    with pytest.raises(im.ArquivoInvalido):
        im.ler("x.json", b"{nao e json")


def test_importar_confere_e_soma_sem_estragar_as_daqui(logged_in_client):
    import io
    from songs_store import get_song, musicas_por_chave, chave_musica, save_song, next_song_id

    with logged_in_client.application.app_context():
        nid = next_song_id()
        save_song({"id": nid, "title": "Importada Daqui", "author": "Igreja",
                   "sections": [{"title": "V", "text": ["versão daqui"]}], "video_url": "https://youtu.be/x"})

    def enviar(*arquivos):
        return logged_in_client.post("/api/importar-musicas/analisar", content_type="multipart/form-data",
                                     data={"arquivos": [(io.BytesIO(d), n) for n, d in arquivos]}).get_json()

    a = enviar(("a.cho", b"{title: importada daqui}\n{artist: IGREJA}\nversao de casa"),
               ("b.cho", b"{title: Importada Nova}\nnova"), ("c.muf", b"?"))
    st = {m["title"]: m["status"] for m in a["musicas"]}
    assert st == {"importada daqui": "diferente", "Importada Nova": "nova"} and len(a["erros"]) == 1
    i_dif = next(m["i"] for m in a["musicas"] if m["status"] == "diferente")
    todas = [m["i"] for m in a["musicas"]]

    # sem marcar "substituir": a daqui fica
    r = logged_in_client.post("/api/importar-musicas/importar", json={"token": a["token"], "importar": todas}).get_json()
    assert (r["novas"], r["puladas"], r["substituidas"]) == (1, 1, 0)
    with logged_in_client.application.app_context():
        assert get_song(nid)["sections"][0]["text"] == ["versão daqui"]

    # substituindo: troca a letra, mantém id e o que só existe aqui (vídeo)
    a = enviar(("a.cho", b"{title: importada daqui}\n{artist: IGREJA}\nversao de casa"))
    r = logged_in_client.post("/api/importar-musicas/importar",
                              json={"token": a["token"], "importar": [0], "substituir": [0]}).get_json()
    assert r["substituidas"] == 1
    with logged_in_client.application.app_context():
        musica = get_song(nid)
        assert musica["sections"][0]["text"] == ["versao de casa"] and musica["video_url"] == "https://youtu.be/x"
        assert chave_musica("Importada Nova", "") in musicas_por_chave()
    assert logged_in_client.post("/api/importar-musicas/importar", json={"token": a["token"]}).status_code == 400


# --- Bíblia em duas versões (blueprints/bible.py: com_segunda_versao) ----

def test_biblia_segunda_versao_vai_junto_pro_telao(app, logged_in_client):
    from blueprints import bible

    payload = {"reference": "João 3:16", "book": "João", "chapter": 3,
               "verses": [{"num": 16, "text": "Porque Deus amou"}], "version": "Bíblia Livre"}
    with app.app_context():
        bible.set_second_translation("")
        assert bible.com_segunda_versao(payload) is payload  # desligado: igual

    assert logged_in_client.post("/api/bible/set_translation2", json={"translation": "xyz"}).status_code == 400
    assert logged_in_client.post("/api/bible/set_translation2", json={"translation": "nbv"}).get_json()["success"]
    assert 'value="nbv" selected' in logged_in_client.get("/bible/jo/3?translation=blivre").get_data(as_text=True)

    operador = _operator_socket(app)
    tela = ucontrol_main.socketio.test_client(app)
    try:
        tela.get_received()
        operador.emit("update_bible", payload)
        recebido = next(m["args"][0] for m in tela.get_received() if m["name"] == "update_bible")
        assert recebido["version2_abbrev"] == "NBV" and recebido["version_abbrev"] == "BLIVRE"
        assert recebido["verses2"][0]["num"] == 16 and "Deus" in recebido["verses2"][0]["text"]
        # tela que abre depois recebe as duas
        assert logged_in_client.get("/api/bible/current").get_json()["version2"] == "Nova Bíblia Viva"
        with app.app_context():
            # a mesma versão nas duas, ou versículo que não existe: só uma
            assert "verses2" not in bible.com_segunda_versao({**payload, "version": "Nova Bíblia Viva"})
            assert "verses2" not in bible.com_segunda_versao({**payload, "verses": [{"num": 999, "text": "x"}]})
    finally:
        operador.disconnect()
        tela.disconnect()
        logged_in_client.post("/api/bible/set_translation2", json={"translation": ""})


# --- Atalhos / Stream Deck (blueprints/atalhos.py) -----------------------

def test_atalhos_exigem_chave_e_comandam_o_telao(app, client, logged_in_client):
    import display_state
    from blueprints import atalhos

    with app.test_request_context():
        chave = atalhos.chave()
    assert chave in logged_in_client.get("/atalhos").get_data(as_text=True)
    assert client.get("/api/atalho/tela-preta").status_code == 401
    assert client.get("/api/atalho/tela-preta?chave=errada").status_code == 401
    assert client.get(f"/api/atalho/nao-existe?chave={chave}").status_code == 404

    tela = ucontrol_main.socketio.test_client(app)
    try:
        tela.get_received()
        r = client.post("/api/atalho/tela-preta-ligar", headers={"X-UControl-Chave": chave})
        assert r.get_json() == {"ok": True, "tela_preta": True} and display_state._blackout_active is True
        assert any(m["name"] == "blackout" and m["args"][0]["active"] for m in tela.get_received())
        assert client.get(f"/api/atalho/tela-preta?chave={chave}").get_json()["tela_preta"] is False

        r = client.get(f"/api/atalho/aviso?chave={chave}&texto=Culto%2019h&telas=public")
        assert r.get_json()["ok"] and client.get("/api/ticker").get_json()["by_target"]["public"] == "Culto 19h"
        assert client.get(f"/api/atalho/aviso?chave={chave}").status_code == 409  # sem texto
        client.get(f"/api/atalho/aviso-tirar?chave={chave}")

        # liberar: o que está preparado vai ao ar
        operador = _operator_socket(app)
        operador.emit("stage_content", {"kind": "image", "payload": {"image_url": "/clock"}, "label": "Relógio"})
        r = client.get(f"/api/atalho/liberar?chave={chave}").get_json()
        assert r == {"ok": True, "liberado": "Relógio"}
        assert client.get(f"/api/atalho/liberar?chave={chave}").get_json()["ok"] is False
        operador.disconnect()

        # sem Controle aberto, "próximo" explica
        r = client.get(f"/api/atalho/proximo?chave={chave}")
        assert r.status_code == 409 and "Controle" in r.get_json()["erro"]
    finally:
        tela.disconnect()
        display_state._blackout_active = False

    with app.test_request_context():
        nova = atalhos.trocar_chave()
    assert nova != chave and client.get(f"/api/atalho/relogio?chave={chave}").status_code == 401


# --- Site no telão (navegador_telao.py, blueprints/site.py) --------------

def test_site_no_telao_so_aparece_com_area_de_trabalho(logged_in_client, monkeypatch):
    import desktop_ponte
    import navegador_telao

    monkeypatch.setattr(desktop_ponte, "_janelas", {})  # Pi/Docker: sem navegador registrado
    html = logged_in_client.get("/control").get_data(as_text=True)
    assert 'data-open-tool="site"' not in html
    r = logged_in_client.post("/api/site/abrir", json={"url": "igreja.com.br"})
    assert r.status_code == 409 and "Windows" in r.get_json()["erro"]

    monkeypatch.setattr(desktop_ponte, "_janelas", {"navegador": lambda: "/bin/chrome", "monitor_telao": lambda: None,
                                                   "perfil_site": "/tmp/x"})
    assert 'data-open-tool="site"' in logged_in_client.get("/control").get_data(as_text=True)
    assert navegador_telao.normalizar("igreja.com.br/agenda") == "https://igreja.com.br/agenda"
    for ruim in ("", "ftp://x", "javascript:alert(1)"):
        with pytest.raises(navegador_telao.SiteErro):
            navegador_telao.normalizar(ruim)


def test_site_no_telao_rotas_e_sites_salvos(logged_in_client, monkeypatch, tmp_path):
    import navegador_telao

    chamadas = []
    monkeypatch.setattr(navegador_telao, "abrir", lambda url: chamadas.append(("abrir", url)) or {"aberto": True, "url": url})
    monkeypatch.setattr(navegador_telao, "comando", lambda a: chamadas.append(("cmd", a)) or {"aberto": True, "zoom": 1.15})
    monkeypatch.setattr(navegador_telao, "fechar", lambda: chamadas.append(("fechar",)))
    assert logged_in_client.post("/api/site/abrir", json={"url": "x.org"}).get_json()["ok"]
    assert logged_in_client.post("/api/site/comando", json={"acao": "zoom-mais"}).get_json()["zoom"] == 1.15
    logged_in_client.post("/api/site/fechar")
    assert chamadas == [("abrir", "x.org"), ("cmd", "zoom-mais"), ("fechar",)]

    monkeypatch.setattr(logged_in_client.application, "instance_path", str(tmp_path))
    r = logged_in_client.post("/api/site/salvos", json={"url": "igreja.com.br", "nome": "Igreja"}).get_json()
    assert r["salvos"] == [{"nome": "Igreja", "url": "https://igreja.com.br"}]
    logged_in_client.post("/api/site/salvos", json={"url": "https://igreja.com.br", "nome": "De novo"})
    assert [s["nome"] for s in logged_in_client.get("/api/site/salvos").get_json()] == ["De novo"]  # sem repetir
    r = logged_in_client.post("/api/site/salvos/remover", json={"url": "https://igreja.com.br"}).get_json()
    assert r["salvos"] == []
    assert logged_in_client.post("/api/site/salvos", json={"url": "ftp://x"}).status_code == 400

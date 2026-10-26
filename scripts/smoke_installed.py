"""Teste rápido do U-Control já instalado e rodando (programa de Windows).
Usado pelo .github/workflows/build-windows.yml depois de instalar o
Setup.exe na máquina do GitHub — pega o que só quebra no .exe empacotado
(arquivo que não entrou no pacote, import escondido, ffmpeg) antes de
alguém baixar o instalador.

    python scripts/smoke_installed.py --exe "C:\\Program Files\\U-Control\\U-Control.exe"

Precisa de um banco novo (sem usuário): cria o primeiro pelo
/primeiro-acesso, como alguém faria na primeira abertura."""
import argparse
import os
import re
import subprocess
import sys
import time

import requests

BASE = 'http://127.0.0.1:8585'


def check(cond, msg):
    print(('ok   ' if cond else 'FALHA') + ' ' + msg, flush=True)
    if not cond:
        sys.exit(1)


def csrf(html):
    m = re.search(r'name="csrf[-_]token"[^>]*(?:value|content)="([^"]+)"', html)
    return m.group(1) if m else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--exe', required=True)
    ap.add_argument('--data-dir', default=os.path.join(os.environ.get('APPDATA', ''), 'U-Control'))
    ap.add_argument('--timeout', type=int, default=90)
    ap.add_argument('--sem-powerpoint', action='store_true', help='instalado sem a tarefa do LibreOffice')
    args = ap.parse_args()
    exe, timeout = args.exe, args.timeout

    s = requests.Session()
    fim = time.time() + timeout
    while True:
        try:
            r = s.get(BASE + '/login', timeout=3)
            break
        except requests.ConnectionError:
            check(time.time() < fim, f'servidor respondeu em {timeout}s')
            time.sleep(1)
    check(r.status_code == 200, 'GET /login')
    check('Repita a senha' in r.text, 'banco novo oferece criar o primeiro usuário')

    r = s.post(BASE + '/primeiro-acesso', allow_redirects=False, data={
        'csrf_token': csrf(r.text), 'username': 'smoke', 'password': 'smoke', 'password_confirm': 'smoke'})
    check(r.status_code == 302 and r.headers['Location'].endswith('/dashboard'), 'primeiro usuário criado')

    r = s.get(BASE + '/control')
    check(r.status_code == 200, 'GET /control logado')
    token = csrf(r.text)

    for url in ('/static/js/telao_bar.js', '/img', '/monitor', '/lyrics/display/public', '/links',
                '/static/qrcodes/control_qr_code.png'):
        check(s.get(BASE + url).status_code == 200, f'GET {url}')

    r = s.get(BASE + '/bible')
    check(r.status_code == 200, 'GET /bible (Bíblias vieram no pacote)')

    r = s.post(BASE + '/add_song', headers={'X-CSRFToken': token}, data={
        'title': 'Smoke', 'author': 'CI', 'sections[]': ['Verso'], 'texts[]': ['Letra de teste']})
    check(r.status_code == 200 and r.json().get('id'), 'música salva')
    r = s.get(BASE + f'/view/{r.json()["id"]}/print')
    check(r.status_code == 200 and r.content.startswith(b'%PDF'), 'PDF de impressão (WeasyPrint)')
    import fitz  # PyMuPDF: o texto saiu mesmo (fonte achada pelo fontconfig)?
    with fitz.open(stream=r.content, filetype='pdf') as pdf:
        check('Letra de teste' in pdf[0].get_text(), 'texto da letra no PDF')

    check(os.path.isfile(os.path.join(args.data_dir, 'instance', 'users.db')), f'banco em {args.data_dir}')

    r = subprocess.run([exe, '-m', 'yt_dlp', '--version'], capture_output=True, text=True, timeout=60)
    check(r.returncode == 0 and r.stdout.strip(), f'yt-dlp pelo .exe ({r.stdout.strip()})')

    ffmpeg = os.path.join(os.path.dirname(exe), 'bin', 'ffmpeg.exe')
    check(os.path.isfile(ffmpeg), f'{ffmpeg} existe')
    r = subprocess.run([ffmpeg, '-version'], capture_output=True, text=True, timeout=30)
    check(r.returncode == 0, 'ffmpeg instalado em bin\\')

    # aviso de atualização: o .exe sabe que é a instalação Windows e a versão
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'VERSION')) as f:
        versao = f.read().strip()
    u = s.get(BASE + '/api/update-status').json()
    check(u['kind'] == 'windows' and u['current'] == versao, f"atualização: windows {u['current']}")

    if not args.sem_powerpoint:
        # PowerPoint -> slides: o instalador baixou e instalou o LibreOffice
        # (tarefa "libreoffice" do ucontrol.iss) e o programa acha ele
        import io
        from pptx import Presentation
        apresentacao = Presentation()
        for texto in ('Slide um', 'Slide dois'):
            apresentacao.slides.add_slide(apresentacao.slide_layouts[1]).shapes.title.text = texto
        # transição no 2º slide: ganha também a versão animada (SVG)
        from lxml import etree
        apresentacao.slides[1]._element.append(etree.fromstring(
            '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:fade/></p:transition>'))
        buf = io.BytesIO()
        apresentacao.save(buf)
        # /upload exige login e token CSRF desde a 1.0.5 (antes era aberto)
        r = s.post(BASE + '/upload', files={'images': ('smoke.pptx', buf.getvalue())},
                   headers={'X-CSRFToken': token})
        check(r.status_code == 200 and 'smoke.pptx' in r.json().get('converting', []), f'PowerPoint aceito pra converter ({r.text[:120]})')
        fim = time.time() + 240
        slides = []
        while time.time() < fim and len(slides) < 2:
            time.sleep(3)
            slides = [u for u in s.get(BASE + '/imagens').json() if 'smoke_slide-' in u]
        check(len(slides) == 2, f'PowerPoint virou {len(slides)} slide(s)')
        fim = time.time() + 240
        animados = []
        while time.time() < fim and len(animados) < 2:
            animados = [f for f in s.get(BASE + '/lista-fotos').json()['fotos'] if f.get('animado')]
            if len(animados) < 2:
                time.sleep(3)
        check(len(animados) == 2, 'PowerPoint com transição ganhou a versão animada')
        if animados:
            svg = s.get(BASE + animados[0]['animado']['svg'])
            check(svg.status_code == 200 and b'aSlideShow' in svg.content, 'SVG animado com o tocador do LibreOffice')

    # e o programa acha esse ffmpeg (PATH montado pelo desktop.py)
    check(s.get(BASE + '/api/system/status').json().get('ffmpeg') is True, 'programa acha o ffmpeg')

    import io
    import zipfile
    r = s.get(BASE + '/api/backup/baixar')
    check(r.status_code == 200 and 'banco/users.db' in zipfile.ZipFile(io.BytesIO(r.content)).namelist(),
          f'backup baixado ({len(r.content) // 1024} KB)')

    # Site no telão: abre no Edge/Chrome do Windows e responde aos comandos
    h = {'X-CSRFToken': csrf(s.get(BASE + '/control').text)}
    r = s.post(BASE + '/api/site/abrir', json={'url': BASE + '/help'}, headers=h)
    check(r.status_code == 200 and r.json().get('aberto'), f'site no telão abre ({r.text[:150]})')
    for acao in ('pagina-baixo', 'rolar-baixo', 'zoom-mais', 'enter', 'recarregar'):
        r = s.post(BASE + '/api/site/comando', json={'acao': acao}, headers=h)
        check(r.status_code == 200 and r.json().get('ok'), f'site no telão: {acao} ({r.text[:120]})')
    r = s.post(BASE + '/api/site/fechar', headers=h)
    check(r.status_code == 200 and not r.json().get('aberto'), 'site no telão fecha')

    # Por último: "Desligar o U-Control" fecha o programa (bandeja inclusive)
    r = s.get(BASE + '/control')
    check('Desligar o U-Control' in r.text, 'menu Sistema tem "Desligar o U-Control"')
    r = s.post(BASE + '/api/desligar', headers={'X-CSRFToken': csrf(r.text)})
    check(r.status_code == 200, 'pedido de desligar aceito')
    nome = os.path.basename(exe)
    fim = time.time() + 20
    while time.time() < fim:
        tarefas = subprocess.run(['tasklist', '/FI', f'IMAGENAME eq {nome}'], capture_output=True, text=True).stdout
        if nome.lower() not in tarefas.lower():
            break
        time.sleep(1)
    check(nome.lower() not in tarefas.lower(), f'{nome} fechou')


if __name__ == '__main__':
    main()
